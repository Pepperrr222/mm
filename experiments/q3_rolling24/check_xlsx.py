from pathlib import Path
import openpyxl
p=Path('experiments/q3_rolling24/results/result3_问题3_结果.xlsx')
wb=openpyxl.load_workbook(p,read_only=True,data_only=False)
print(wb.sheetnames)
for ws in wb.worksheets: print(ws.title,ws.max_row,ws.max_column,list(ws.values)[:2])