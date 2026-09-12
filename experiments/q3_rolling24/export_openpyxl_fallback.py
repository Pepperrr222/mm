import json, shutil
from pathlib import Path
import openpyxl
root=Path(__file__).resolve().parents[1]
out=root/'results'; data=json.loads((out/'workbook_data.json').read_text(encoding='utf-8'))
template=root.parent/'01_题目与数据'/'原始附件'/'附件5'/'result3.xlsx'
dest=out/'result3_问题3_结果.xlsx'; shutil.copy2(template,dest)
wb=openpyxl.load_workbook(dest)
headers=data['headers']
for name,rows in [('计划购电量',data['plan']),('调整购电量',data['final'])]:
    ws=wb[name]
    for j,v in enumerate(headers,1): ws.cell(1,j).value=v
    for i,row in enumerate(rows,2):
        for j,v in enumerate(row,1): ws.cell(i,j).value=v
    for i in range(2,337): ws.cell(i,1).number_format='yyyy-mm-dd'
    for row in ws.iter_rows(min_row=2,max_row=335,min_col=2,max_col=166):
        for c in row: c.number_format='#,##0.000000'
    for i in range(2,336): ws.cell(i,171).value=f'=SUM(B{i}:EO{i})'
ws=wb['充放电量'];ws.delete_rows(2,ws.max_row)
for i,row in enumerate(data['storage'],2):
    for j,v in enumerate(row,1): ws.cell(i,j).value=v
for i in range(2,2+len(data['storage'])):
    ws.cell(i,1).number_format='yyyy-mm-dd'
    for j in (3,4,6): ws.cell(i,j).number_format='#,##0.000000'
ws=wb['紧急购电量'];ws.delete_rows(1,ws.max_row);ws.append(['日期','紧急时段','购电量'])
for row in data['emergency']: ws.append(row)
for i in range(2,2+len(data['emergency'])):
    ws.cell(i,1).number_format='yyyy-mm-dd';ws.cell(i,3).number_format='#,##0.000000'
wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True
wb.save(dest)
print(dest)