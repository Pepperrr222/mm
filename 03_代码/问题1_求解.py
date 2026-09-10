from __future__ import annotations

import hashlib
import importlib.metadata
import json
import platform
import shutil
import sys
from datetime import datetime, time
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import lil_matrix

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "01_题目与数据" / "原始附件" / "附件1.xlsx"
TEMPLATE = ROOT / "01_题目与数据" / "原始附件" / "附件5" / "result1.xlsx"
RESULTS = ROOT / "04_结果"
FIGURES = ROOT / "05_图表"
RESULTS.mkdir(exist_ok=True)
FIGURES.mkdir(exist_ok=True)

N, DT = 144, 1 / 6
ETA_C = ETA_D = 0.9
E_MIN, E_MAX, E0 = 1200.0, 10800.0, 6000.0
P_MAX = 5000.0
Q_MAX = P_MAX * DT
TAU = 1e-6


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def label(t: int) -> str:
    a = t * 10
    b = (t + 1) * 10
    def fmt(minutes: int) -> str:
        if minutes == 1440:
            return "0:00+1"
        return f"{minutes // 60}:{minutes % 60:02d}"
    return f"{fmt(a)}-{fmt(b)}"


def read_input() -> pd.DataFrame:
    wb = openpyxl.load_workbook(INPUT, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == N + 1 and len(rows[0]) == 4
    vals = []
    for i, row in enumerate(rows[1:]):
        assert all(x is not None for x in row), f"附件1第{i+2}行存在空值"
        vals.append(row)
    df = pd.DataFrame(vals, columns=["time_raw", "price", "load_kw", "pv_kw"])
    assert len(df) == N
    # 时间字段允许 Excel time 和字符串混合；索引以数据行号为准，并核对右端点序列。
    expected_minutes = [(i + 1) * 10 for i in range(N)]
    for i, value in enumerate(df["time_raw"]):
        if isinstance(value, time):
            minutes = value.hour * 60 + value.minute
        else:
            text = str(value).strip().replace("+1", "")
            hh, mm = text.split(":")[:2]
            minutes = int(hh) * 60 + int(mm)
        if minutes == 0 and i == N - 1:
            minutes = 1440
        assert minutes == expected_minutes[i], (i, value, minutes, expected_minutes[i])
    for c in ["price", "load_kw", "pv_kw"]:
        df[c] = pd.to_numeric(df[c], errors="raise")
        assert np.isfinite(df[c]).all()
    return df


def idx(kind: int, t: int) -> int:
    return kind * N + t


def solve(df: pd.DataFrame):
    price = df.price.to_numpy(float)
    load = df.load_kw.to_numpy(float) * DT
    pv = df.pv_kw.to_numpy(float) * DT
    # variables: g, c, d, w (0..N-1), E (0..N), z (0..N-1)
    nv = 5 * N + (N + 1)
    ebase, zbase = 4 * N, 5 * N + 1
    c1 = np.zeros(nv)
    c1[:N] = price
    lb = np.zeros(nv)
    ub = np.full(nv, np.inf)
    ub[N:3*N] = Q_MAX
    lb[ebase:ebase+N+1] = E_MIN
    ub[ebase:ebase+N+1] = E_MAX
    lb[zbase:] = 0
    ub[zbase:] = 1
    integrality = np.zeros(nv)
    integrality[zbase:] = 1
    A = lil_matrix((4*N + 2, nv), dtype=float)
    b_l = np.full(4*N + 2, -np.inf)
    b_u = np.full(4*N + 2, np.inf)
    r = 0
    for t in range(N):
        # g + d - c - w = load - pv
        A[r, idx(0,t)] = 1; A[r, idx(1,t)] = -1; A[r, idx(2,t)] = 1; A[r, idx(3,t)] = -1
        b_l[r] = b_u[r] = load[t] - pv[t]; r += 1
        # E[t+1] - E[t] - eta_c*c + d/eta_d = 0
        A[r, ebase+t+1] = 1; A[r, ebase+t] = -1; A[r, idx(1,t)] = -ETA_C; A[r, idx(2,t)] = 1/ETA_D
        b_l[r] = b_u[r] = 0; r += 1
        # c <= Q_MAX z, d <= Q_MAX(1-z)
        A[r, idx(1,t)] = 1; A[r, zbase+t] = -Q_MAX; b_u[r] = 0; r += 1
        A[r, idx(2,t)] = 1; A[r, zbase+t] = Q_MAX; b_u[r] = Q_MAX; r += 1
    # E0=E144=6000
    A[r, ebase] = 1; b_l[r] = b_u[r] = E0; r += 1
    A[r, ebase+N] = 1; b_l[r] = b_u[r] = E0; r += 1
    A = A.tocsr()
    constraints = [LinearConstraint(A, b_l, b_u)]
    res1 = milp(c1, integrality=integrality, bounds=Bounds(lb, ub), constraints=constraints,
                options={"mip_rel_gap": 0.0})
    if not res1.success:
        raise RuntimeError(res1.message)
    cstar = float(price @ res1.x[:N])
    # Stage 2: throughput with explicit cost tolerance.
    A2 = lil_matrix((A.shape[0] + 1, nv), dtype=float)
    A2[:-1] = A
    A2[-1, :N] = price
    lo2 = np.r_[b_l, -np.inf]
    hi2 = np.r_[b_u, cstar + TAU]
    c2 = np.zeros(nv); c2[N:3*N] = 1
    res2 = milp(c2, integrality=integrality, bounds=Bounds(lb, ub),
                constraints=[LinearConstraint(A2.tocsr(), lo2, hi2)], options={"mip_rel_gap": 0.0})
    if not res2.success:
        raise RuntimeError(res2.message)
    x = res2.x
    out = df.copy()
    out["interval"] = [label(t) for t in range(N)]
    out["grid_kwh"] = x[:N]
    out["charge_kwh"] = x[N:2*N]
    out["discharge_kwh"] = x[2*N:3*N]
    out["curtail_kwh"] = x[3*N:4*N]
    out["E_start_kwh"] = x[ebase:ebase+N]
    out["E_end_kwh"] = x[ebase+1:ebase+N+1]
    out["z_charge"] = np.rint(x[zbase:]).astype(int)
    balance = out.grid_kwh + out.pv_kw * DT + out.discharge_kwh - out.load_kw * DT - out.charge_kwh - out.curtail_kwh
    state = out.E_end_kwh - out.E_start_kwh - ETA_C*out.charge_kwh + out.discharge_kwh/ETA_D
    metrics = {
        "stage1_cost_yuan": cstar,
        "stage2_cost_yuan": float(price @ x[:N]),
        "cost_gap_yuan": float(price @ x[:N] - cstar),
        "grid_kwh": float(x[:N].sum()),
        "charge_kwh": float(x[N:2*N].sum()),
        "discharge_kwh": float(x[2*N:3*N].sum()),
        "curtail_kwh": float(x[3*N:4*N].sum()),
        "E_min_kwh": float(x[ebase:ebase+N+1].min()),
        "E_max_kwh": float(x[ebase:ebase+N+1].max()),
        "max_balance_residual_kwh": float(np.max(np.abs(balance))),
        "max_state_residual_kwh": float(np.max(np.abs(state))),
        "simultaneous_periods": int(np.sum((out.charge_kwh > 1e-7) & (out.discharge_kwh > 1e-7))),
        "throughput_kwh": float((x[N:3*N]).sum()),
    }
    if metrics["cost_gap_yuan"] > TAU + 1e-9:
        raise RuntimeError(f"二阶段费用超出容差: {metrics['cost_gap_yuan']}")
    return out, metrics, cstar, price


def write_tables(out: pd.DataFrame, metrics: dict, price: np.ndarray):
    out.to_csv(RESULTS / "问题1_逐时段结果.csv", index=False, encoding="utf-8-sig")
    four = []
    for k in range(6):
        part = out.iloc[k*24:(k+1)*24]
        four.append({"interval": f"{k*4}:00-{(k+1)*4}:00", "charge_kwh": part.charge_kwh.sum(), "discharge_kwh": part.discharge_kwh.sum()})
    pd.DataFrame(four).to_csv(RESULTS / "问题1_四小时汇总.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame([metrics]).to_csv(RESULTS / "问题1_指标.csv", index=False, encoding="utf-8-sig")
    # Preserve template sheets; fill only cells specified in the model contract.
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb["计划购电量"]
    for t in range(N):
        ws.cell(t+2, 1).value = out.iloc[t].interval
        ws.cell(t+2, 2).value = float(out.iloc[t].grid_kwh)
    ws2 = wb["充放电量"]
    for k in range(6):
        part = out.iloc[k*24:(k+1)*24]
        ws2.cell(k+2, 1).value = f"{k*4}:00-{(k+1)*4}:00"
        ws2.cell(k+2, 2).value = float(part.charge_kwh.sum())
        ws2.cell(k+2, 3).value = float(part.discharge_kwh.sum())
    ws2["D2"] = "0:00"; ws2["D3"] = "24:00"
    ws2["E2"] = E0; ws2["E3"] = E0
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(RESULTS / "result1_问题1_结果.xlsx")


def save_fig(fig, name):
    sys.path.insert(0, r"C:\Users\Lenovo\.codex\skills\math-modeling\tools\figure\scripts")
    from export_figure import export_figure
    paths = export_figure(fig, str(FIGURES / name), formats=["svg", "png"], dpi=300, size_inches=(6.4, 3.8), grayscale_preview=True)
    preview_dir = FIGURES / "previews"
    preview_dir.mkdir(exist_ok=True)
    for p in paths:
        if str(p).endswith("_grayscale.png"):
            shutil.move(str(p), preview_dir / Path(p).name)


def make_figures(df, out, metrics):
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.size": 9, "axes.titlesize": 10, "axes.labelsize": 9})
    x = np.arange(N) / 6
    # raw: trend, energy profile, distribution/relationship
    fig, ax = plt.subplots(); ax.plot(x, df.price, label="price"); ax.set(xlabel="hour", ylabel="yuan/kWh", title="Raw electricity price"); ax.grid(alpha=.2); save_fig(fig, "raw_q1_price"); plt.close(fig)
    fig, ax = plt.subplots(); ax.plot(x, df.load_kw, label="load"); ax.plot(x, df.pv_kw, label="PV"); ax.set(xlabel="hour", ylabel="kW", title="Raw load and PV forecast"); ax.legend(); ax.grid(alpha=.2); save_fig(fig, "raw_q1_load_pv"); plt.close(fig)
    fig, ax = plt.subplots(); ax.scatter(df.load_kw, df.pv_kw, s=8, alpha=.6); ax.set(xlabel="load / kW", ylabel="PV / kW", title="Load–PV relationship"); ax.grid(alpha=.2); save_fig(fig, "raw_q1_load_pv_scatter"); plt.close(fig)
    # process
    net = df.load_kw.to_numpy()*DT - df.pv_kw.to_numpy()*DT
    fig, ax = plt.subplots(); ax.plot(x, net, color="#555"); ax.axhline(0, color="k", lw=.7); ax.set(xlabel="hour", ylabel="net energy / kWh", title="Net-load process input"); ax.grid(alpha=.2); save_fig(fig, "process_q1_netload"); plt.close(fig)
    fig, ax = plt.subplots(); ax.plot(x, out.E_start_kwh, color="#0072B2"); ax.axhline(E_MIN, ls="--", color="gray"); ax.axhline(E_MAX, ls="--", color="gray"); ax.set(xlabel="hour", ylabel="stored energy / kWh", title="Storage-state trajectory"); ax.grid(alpha=.2); save_fig(fig, "process_q1_storage"); plt.close(fig)
    fig, ax = plt.subplots(); ax.plot(x, out.grid_kwh, label="grid"); ax.plot(x, out.charge_kwh, label="charge"); ax.plot(x, out.discharge_kwh, label="discharge"); ax.set(xlabel="hour", ylabel="energy / kWh", title="Dispatch process"); ax.legend(ncol=3); ax.grid(alpha=.2); save_fig(fig, "process_q1_dispatch"); plt.close(fig)
    # result
    fig, ax = plt.subplots(); ax.step(x, out.grid_kwh, where="post", color="#D55E00"); ax.set(xlabel="hour", ylabel="grid purchase / kWh", title="Optimal grid-purchase schedule"); ax.grid(alpha=.2); save_fig(fig, "result_q1_grid"); plt.close(fig)
    baseline = float(np.sum(np.maximum(net, 0) * df.price.to_numpy()))
    fig, ax = plt.subplots(); ax.bar(["without storage", "optimized"], [baseline, metrics["stage2_cost_yuan"]], color=["#999999", "#0072B2"]); ax.set(ylabel="cost / yuan", title="Cost reduction from storage"); ax.grid(axis="y", alpha=.2); save_fig(fig, "result_q1_cost"); plt.close(fig)
    f = out.groupby(out.index // 24)[["charge_kwh", "discharge_kwh"]].sum(); fig, ax = plt.subplots(); f.plot.bar(ax=ax, color=["#009E73", "#CC79A7"]); ax.set(xlabel="four-hour block", ylabel="energy / kWh", title="Four-hour charge/discharge totals"); ax.legend(); ax.grid(axis="y", alpha=.2); save_fig(fig, "result_q1_four_hour"); plt.close(fig)


def main():
    df = read_input()
    out, metrics, cstar, price = solve(df)
    write_tables(out, metrics, price)
    df.assign(interval=[label(t) for t in range(N)], net_load_kwh=df.load_kw*DT-df.pv_kw*DT).to_csv(RESULTS / "问题1_输入审计.csv", index=False, encoding="utf-8-sig")
    df[["price", "load_kw", "pv_kw"]].assign(net_load_kwh=df.load_kw*DT-df.pv_kw*DT).to_csv(RESULTS / "问题1_数值剖析.csv", index=False, encoding="utf-8-sig")
    make_figures(df, out, metrics)
    manifest = {
        "generated_at": datetime.now().isoformat(), "python": sys.version, "platform": platform.platform(),
        "command": "python 03_代码/问题1_求解.py", "seed": None,
        "inputs": {str(INPUT.relative_to(ROOT)): sha256(INPUT), str(TEMPLATE.relative_to(ROOT)): sha256(TEMPLATE)},
        "code_sha256": sha256(Path(__file__)),
        "outputs": {str(p.relative_to(ROOT)): sha256(p) for p in [RESULTS / "问题1_逐时段结果.csv", RESULTS / "问题1_四小时汇总.csv", RESULTS / "问题1_指标.csv", RESULTS / "result1_问题1_结果.xlsx"]},
        "dependencies": {name: importlib.metadata.version(name) for name in ["numpy", "pandas", "scipy", "openpyxl", "matplotlib"]},
        "parameters": {"N": N, "dt_h": DT, "eta_c": ETA_C, "eta_d": ETA_D, "E_min_kwh": E_MIN, "E_max_kwh": E_MAX, "E0_kwh": E0, "P_max_kw": P_MAX, "tau_yuan": TAU},
        "metrics": metrics,
    }
    (RESULTS / "复现清单.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(metrics, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
