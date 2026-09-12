"""第四问独立实现：论文复现框架 + Q3 最近2小时负荷偏差修正。

输出与既有第四问完全隔离。计划阶段只使用决策时点之前可见的信息；
实际负荷、光伏和附件4实际电价只用于事后执行与结算。
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import platform
import shutil
import sys
import time
import types
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
PAPER = ROOT / "paper_reproduction_package"
IMPROVEMENT = ROOT / "paper_q3_improvements" / "load_bias_2h"
ATTACHMENT4 = ROOT / "01_题目与数据" / "原始附件" / "附件4.xlsx"
TEMPLATE_DIR = ROOT / "01_题目与数据" / "原始附件" / "附件5"
DEFAULT_OUT = ROOT / "04_结果" / "问题4_队友Q3优化"
N = 144
ISSUES = (0, 36, 72, 108)
DT = 1.0 / 6.0
SEED = 20260912


def _load_paper_modules():
    """加载论文复现代码，同时绕开其历史目录定位错误，不修改队友文件。"""
    q2_path = ROOT / "03_代码" / "问题2_求解.py"
    spec = importlib.util.spec_from_file_location("q4_team_q2", q2_path)
    q2 = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = q2
    assert spec.loader is not None
    spec.loader.exec_module(q2)

    shim = types.ModuleType("q2_forecast")
    shim.q2 = q2
    shim.day_type = lambda ts: 0 if pd.Timestamp(ts).weekday() in (4, 5) else 1
    sys.modules["q2_forecast"] = shim
    sys.path.insert(0, str(PAPER))
    import q3_reproduce as p3
    return q2, p3


q2, p3 = _load_paper_modules()
E0, EMIN, EMAX, ETA, S = p3.E0, p3.EMIN, p3.EMAX, p3.ETA, p3.S


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def read_price4(dates: pd.DatetimeIndex) -> np.ndarray:
    raw = pd.read_excel(ATTACHMENT4, header=None)
    price_dates = pd.to_datetime(raw.iloc[1:, 0], errors="coerce")
    values = raw.iloc[1:, 1:].apply(pd.to_numeric, errors="coerce").to_numpy(float)
    if values.shape != (365, N) or price_dates.isna().any() or not np.isfinite(values).all():
        raise ValueError(f"附件4结构异常: dates={len(price_dates)}, values={values.shape}")
    if not np.array_equal(price_dates.to_numpy("datetime64[ns]"), dates.to_numpy("datetime64[ns]")):
        raise ValueError("附件4日期与负荷/光伏日期不一致")
    if (values <= 0).any():
        raise ValueError("附件4存在非正电价")
    return values


def prepare_forecasts():
    data, base = p3.prepare()
    corrected_path = IMPROVEMENT / "load_bias_cache.npz"
    daily = np.load(PAPER / "cache_q2_forecasts.npz")["load_kw"]
    corrected = base["load"].copy()
    for day in range(1, 365):
        for issue_no, k in enumerate(ISSUES):
            if k == 0:
                continue
            bias = float(np.median(data.load_kw[day, k-12:k] - daily[day, k-12:k]))
            corrected[day, issue_no] = np.maximum(
                0.0, base["load"][day, issue_no] + np.exp(-np.arange(N) / 36.0) * bias
            )
    # 队友缓存只作为一致性证据，不再是运行依赖。
    if corrected_path.exists():
        cached = np.load(corrected_path)["load"]
        if cached.shape != corrected.shape or not np.allclose(
                cached, corrected, atol=1e-9, rtol=0.0, equal_nan=True):
            raise ValueError("代码内重建的2小时偏差修正与队友缓存不一致")
    return data, base, corrected


def scenario_set(day: int, issue_no: int, base: dict, corrected: np.ndarray, use_bias: bool):
    candidates = [i for i in range(max(1, day - 35), day)
                  if np.isfinite(base["actual_load"][i, issue_no]).all()][-30:]
    if not candidates:
        raise ValueError(f"day={day}, issue={issue_no}: 无历史残差场景")
    idx = np.asarray(candidates)
    forecast_load = corrected if use_bias else base["load"]
    load_error = base["actual_load"][idx, issue_no] - forecast_load[idx, issue_no]
    pv_error = base["actual_pv"][idx, issue_no] - base["pv"][idx, issue_no]
    load_s = np.maximum(0.0, forecast_load[day, issue_no] + load_error)
    pv_s = np.maximum(0.0, base["pv"][day, issue_no] + pv_error)
    return load_s, pv_s, len(idx)


def causal_price_clock(price4: np.ndarray, day: int, issue_slot: int) -> np.ndarray:
    """构造 solve_roll 所需的钟点价格：采用7日前同槽价格，杜绝未来泄露。"""
    clock = np.empty(N)
    for z in range(N):
        absolute_slot = issue_slot + z
        target_day = min(day + absolute_slot // N, 364)
        slot = absolute_slot % N
        history_day = max(0, target_day - 7)
        clock[slot] = price4[history_day, slot]
    return clock


def execute_block(data, price4, day, start, stop, plan, reserve, soc,
                  delivered, emergency, waste, charge, discharge, soc_start, soc_end,
                  issue_slot, price_forecast):
    issue_price = causal_price_clock(price4, day, issue_slot)
    for t in range(start, stop):
        z = t - start
        e_before = soc
        gap = (data.load_kw[day, t] - data.pv_kw[day, t]) * DT - plan[z]
        d = min(gap, S, ETA * max(soc - reserve[z], 0.0)) if gap > 0 else 0.0
        c = min(-gap, S, (EMAX - soc) / ETA) if gap <= 0 else 0.0
        h = max(gap + c - d, 0.0)
        w = max(-gap - c + d, 0.0)
        soc = soc + ETA * c - d / ETA
        if not EMIN - 1e-7 <= soc <= EMAX + 1e-7:
            raise AssertionError((day, t, soc))
        delivered[t], emergency[t], waste[t] = plan[z], h, w
        charge[t], discharge[t] = c, d
        soc_start[t], soc_end[t] = e_before, soc
        price_forecast[t] = issue_price[t]
    return soc


def simulate(part: str, start_day: int, end_day: int, use_bias: bool):
    data, base, corrected = prepare_forecasts()
    dates = pd.DatetimeIndex(data.dates)
    price4 = read_price4(dates)
    updates = (0,) if part == "q42" else ISSUES
    soc = float(E0)
    daily_rows, detail_rows = [], []
    started = time.perf_counter()
    for day in range(start_day, end_day):
        g0 = None
        current_commitment = None
        delivered = np.zeros(N); emergency = np.zeros(N); waste = np.zeros(N)
        charge = np.zeros(N); discharge = np.zeros(N)
        step_up = np.zeros(N); step_down = np.zeros(N)
        soc_start = np.zeros(N); soc_end = np.zeros(N); forecast_price = np.zeros(N)
        scenario_counts = []
        initial_soc = soc
        for pos, k in enumerate(updates):
            issue_no = ISSUES.index(k)
            load_s, pv_s, m = scenario_set(day, issue_no, base, corrected, use_bias and k > 0)
            scenario_counts.append(m)
            price_clock = causal_price_clock(price4, day, k)
            plan = p3.solve_roll(price_clock, current_commitment, k, load_s, pv_s, soc)
            if k == 0:
                g0 = plan.copy()
                current_commitment = plan.copy()
            else:
                previous = current_commitment[k:].copy()
                step_up[k:] += np.maximum(plan - previous, 0.0)
                step_down[k:] += np.maximum(previous - plan, 0.0)
                current_commitment[k:] = plan
            stop = updates[pos + 1] if pos + 1 < len(updates) else N
            reserve = p3.reserves(price_clock, (load_s - pv_s) * DT, plan, N - k)
            soc = execute_block(data, price4, day, k, stop, plan, reserve, soc,
                                delivered, emergency, waste, charge, discharge,
                                soc_start, soc_end, k, forecast_price)
        assert g0 is not None
        actual_price = price4[day]
        baseline_cost = float(actual_price @ g0)
        up_cost = float(1.5 * actual_price @ step_up)
        down_refund = float(0.5 * actual_price @ step_down)
        normal_cost = baseline_cost + up_cost - down_refund
        emergency_cost = float(5.0 * actual_price @ emergency)
        total_cost = normal_cost + emergency_cost
        daily_rows.append({
            "date": dates[day], "part": part, "variant": "P1_bias2h" if use_bias else "P0",
            "initial_soc_kwh": initial_soc, "final_soc_kwh": soc,
            "baseline_plan_cost_yuan": baseline_cost,
            "upward_adjustment_cost_yuan": up_cost,
            "downward_cancellation_refund_yuan": down_refund,
            "settled_normal_cost_yuan": normal_cost,
            "emergency_cost_yuan": emergency_cost, "total_cost_yuan": total_cost,
            "original_plan_kwh": float(g0.sum()), "final_commitment_kwh": float(delivered.sum()),
            "up_kwh": float(step_up.sum()), "down_kwh": float(step_down.sum()),
            "emergency_kwh": float(emergency.sum()), "waste_kwh": float(waste.sum()),
            "charge_kwh": float(charge.sum()), "discharge_kwh": float(discharge.sum()),
            "scenario_count_min": int(min(scenario_counts)),
        })
        for t in range(N):
            balance = delivered[t] + discharge[t] + emergency[t] + data.pv_kw[day, t] * DT \
                      - data.load_kw[day, t] * DT - charge[t] - waste[t]
            state = soc_end[t] - soc_start[t] - ETA * charge[t] + discharge[t] / ETA
            detail_rows.append({
                "date": dates[day], "slot": t + 1, "time_start": f"{t//6:02d}:{(t%6)*10:02d}",
                "actual_price_yuan_per_kwh": actual_price[t],
                "forecast_price_yuan_per_kwh": forecast_price[t],
                "load_kw": data.load_kw[day, t], "pv_kw": data.pv_kw[day, t],
                "original_plan_kwh": g0[t], "final_commitment_kwh": delivered[t],
                "adjustment_kwh": delivered[t] - g0[t],
                "up_adjustment_kwh": step_up[t], "down_cancellation_kwh": step_down[t],
                "charge_kwh": charge[t],
                "discharge_kwh": discharge[t], "emergency_kwh": emergency[t], "waste_kwh": waste[t],
                "soc_start_kwh": soc_start[t], "soc_end_kwh": soc_end[t],
                "balance_residual_kwh": balance, "state_residual_kwh": state,
            })
        print(f"[{part}] {dates[day].date()} total={total_cost:.2f} soc={soc:.2f}", flush=True)
    daily = pd.DataFrame(daily_rows)
    detail = pd.DataFrame(detail_rows)
    return data, price4, daily, detail, time.perf_counter() - started


def emergency_intervals(group: pd.DataFrame, tol: float = 1e-8):
    slots = group.loc[group.emergency_kwh > tol, "slot"].astype(int).to_list()
    if not slots:
        return []
    chunks, current = [], [slots[0]]
    for x in slots[1:]:
        if x == current[-1] + 1:
            current.append(x)
        else:
            chunks.append(current); current = [x]
    chunks.append(current)
    out = []
    for c in chunks:
        a, b = c[0] - 1, c[-1]
        end_label = "00:00+1" if b == N else f"{(b//6)%24:02d}:{(b%6)*10:02d}"
        label = f"{a//6:02d}:{(a%6)*10:02d}-{end_label}"
        energy = float(group.loc[group.slot.isin(c), "emergency_kwh"].sum())
        out.append((label, energy))
    return out


def write_official_xlsx(part: str, daily: pd.DataFrame, detail: pd.DataFrame, out: Path):
    template = TEMPLATE_DIR / ("result4-2.xlsx" if part == "q42" else "result4-3.xlsx")
    target = out / template.name
    shutil.copy2(template, target)
    wb = load_workbook(target)
    plan_ws = wb.worksheets[0]
    adjusted_ws = wb.worksheets[1] if part == "q43" else None
    storage_ws = wb.worksheets[2] if part == "q43" else wb.worksheets[1]
    emergency_ws = wb.worksheets[3] if part == "q43" else wb.worksheets[2]
    by_date = {pd.Timestamp(d).normalize(): g.sort_values("slot") for d, g in detail.groupby("date")}
    daily_idx = daily.set_index(pd.to_datetime(daily.date).dt.normalize())
    for row in range(2, plan_ws.max_row + 1):
        date = pd.Timestamp(plan_ws.cell(row, 1).value).normalize()
        if date not in by_date:
            continue
        g = by_date[date]
        for t, value in enumerate(g.original_plan_kwh, start=2):
            plan_ws.cell(row, t, float(value))
        plan_ws.cell(row, 146, float(g.original_plan_kwh.sum()))
        plan_ws.cell(row, 147, float((g.actual_price_yuan_per_kwh * g.original_plan_kwh).sum()))
        if adjusted_ws is not None:
            for t, value in enumerate(g.final_commitment_kwh, start=2):
                adjusted_ws.cell(row, t, float(value))
            adjusted_ws.cell(row, 146, float(g.final_commitment_kwh.sum()))
            settlement = float(daily_idx.loc[date, "settled_normal_cost_yuan"])
            adjusted_ws.cell(row, 147, settlement)
    storage_ws.delete_rows(2, max(storage_ws.max_row - 1, 0))
    emergency_ws.delete_rows(2, max(emergency_ws.max_row - 1, 0))
    for date, g in by_date.items():
        for block in range(6):
            lo, hi = block * 24, (block + 1) * 24
            row = storage_ws.max_row + 1
            storage_ws.cell(row, 1, date.to_pydatetime() if block == 0 else None)
            storage_ws.cell(row, 2, f"{4*block}:00-{4*(block+1)}:00")
            storage_ws.cell(row, 3, float(g.iloc[lo:hi].charge_kwh.sum()))
            storage_ws.cell(row, 4, float(g.iloc[lo:hi].discharge_kwh.sum()))
            if block == 0:
                storage_ws.cell(row, 5, "0:00")
                storage_ws.cell(row, 6, float(g.iloc[0].soc_start_kwh))
            elif block == 1:
                storage_ws.cell(row, 5, "24:00")
                storage_ws.cell(row, 6, float(g.iloc[-1].soc_end_kwh))
        intervals = emergency_intervals(g)
        if not intervals:
            emergency_ws.append([date.to_pydatetime(), "无", 0.0])
        else:
            for i, (label, energy) in enumerate(intervals):
                emergency_ws.append([date.to_pydatetime() if i == 0 else None, label, energy])
    wb.save(target)
    return target


def write_outputs(part: str, daily: pd.DataFrame, detail: pd.DataFrame, runtime: float,
                  out_root: Path, args):
    out = out_root / ("问题4_2" if part == "q42" else "问题4_3")
    out.mkdir(parents=True, exist_ok=True)
    daily.to_csv(out / "每日指标.csv", index=False, encoding="utf-8-sig")
    detail.to_csv(out / "逐时段结果.csv", index=False, encoding="utf-8-sig")
    specified = pd.to_datetime(["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"])
    daily[daily.date.isin(specified)].to_csv(out / "指定日期结果.csv", index=False, encoding="utf-8-sig")
    expected_days = args.end_day - args.start_day
    continuity = 0.0 if len(daily) <= 1 else float(np.max(np.abs(
        daily.initial_soc_kwh.to_numpy()[1:] - daily.final_soc_kwh.to_numpy()[:-1])))
    baseline_recalc = detail.groupby("date").apply(
        lambda g: float((g.actual_price_yuan_per_kwh * g.original_plan_kwh).sum()),
        include_groups=False).to_numpy()
    normal_recalc = detail.groupby("date").apply(
        lambda g: float((g.actual_price_yuan_per_kwh *
                         (g.original_plan_kwh + 1.5*g.up_adjustment_kwh -
                          0.5*g.down_cancellation_kwh)).sum()), include_groups=False).to_numpy()
    emergency_recalc = detail.groupby("date").apply(
        lambda g: float((5*g.actual_price_yuan_per_kwh*g.emergency_kwh).sum()),
        include_groups=False).to_numpy()
    cost_residual = float(max(np.max(np.abs(baseline_recalc-daily.baseline_plan_cost_yuan)),
                              np.max(np.abs(normal_recalc-daily.settled_normal_cost_yuan)),
                              np.max(np.abs(emergency_recalc-daily.emergency_cost_yuan))))
    simultaneous_cd = int(((detail.charge_kwh > 1e-8) & (detail.discharge_kwh > 1e-8)).sum())
    emergency_charge = int(((detail.emergency_kwh > 1e-8) & (detail.charge_kwh > 1e-8)).sum())
    emergency_waste = int(((detail.emergency_kwh > 1e-8) & (detail.waste_kwh > 1e-8)).sum())
    passed = bool(len(daily) == expected_days and len(detail) == expected_days*N and
                  detail.balance_residual_kwh.abs().max() < 1e-7 and
                  detail.state_residual_kwh.abs().max() < 1e-7 and continuity < 1e-7 and
                  detail.soc_start_kwh.min() >= EMIN-1e-7 and detail.soc_end_kwh.max() <= EMAX+1e-7 and
                  detail.charge_kwh.max() <= S+1e-7 and detail.discharge_kwh.max() <= S+1e-7 and
                  simultaneous_cd == 0 and emergency_charge == 0 and emergency_waste == 0 and
                  cost_residual < 1e-6)
    validation = {
        "passed": passed,
        "days": len(daily), "detail_rows": len(detail),
        "max_balance_residual_kwh": float(detail.balance_residual_kwh.abs().max()),
        "max_state_residual_kwh": float(detail.state_residual_kwh.abs().max()),
        "storage_min_kwh": float(min(detail.soc_start_kwh.min(), detail.soc_end_kwh.min())),
        "storage_max_kwh": float(max(detail.soc_start_kwh.max(), detail.soc_end_kwh.max())),
        "max_charge_kwh_per_slot": float(detail.charge_kwh.max()),
        "max_discharge_kwh_per_slot": float(detail.discharge_kwh.max()),
        "simultaneous_charge_discharge": simultaneous_cd,
        "simultaneous_emergency_charge": emergency_charge,
        "simultaneous_emergency_waste": emergency_waste,
        "max_cross_day_soc_residual_kwh": continuity,
        "max_cost_recalculation_residual_yuan": cost_residual,
        "total_cost_yuan": float(daily.total_cost_yuan.sum()),
        "emergency_cost_yuan": float(daily.emergency_cost_yuan.sum()),
    }
    (out / "总体验证.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8")
    xlsx = write_official_xlsx(part, daily, detail, out)
    wb = load_workbook(xlsx, read_only=True, data_only=False)
    validation["xlsx_sheet_count"] = len(wb.sheetnames)
    validation["xlsx_expected_sheet_count"] = 3 if part == "q42" else 4
    validation["xlsx_structure_ok"] = validation["xlsx_sheet_count"] == validation["xlsx_expected_sheet_count"]
    validation["passed"] = bool(validation["passed"] and validation["xlsx_structure_ok"])
    (out / "总体验证.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8")
    input_hashes = {
        str(ATTACHMENT4): sha256(ATTACHMENT4),
        str(ROOT / "01_题目与数据" / "原始附件" / "附件2.xlsx"): sha256(ROOT / "01_题目与数据" / "原始附件" / "附件2.xlsx"),
        str(ROOT / "01_题目与数据" / "原始附件" / "附件3.xlsx"): sha256(ROOT / "01_题目与数据" / "原始附件" / "附件3.xlsx"),
        str(TEMPLATE_DIR / ("result4-2.xlsx" if part == "q42" else "result4-3.xlsx")): sha256(TEMPLATE_DIR / ("result4-2.xlsx" if part == "q42" else "result4-3.xlsx")),
        str(PAPER / "q3_reproduce.py"): sha256(PAPER / "q3_reproduce.py"),
        str(PAPER / "q2_reproduce.py"): sha256(PAPER / "q2_reproduce.py"),
        str(PAPER / "cache_q2_forecasts.npz"): sha256(PAPER / "cache_q2_forecasts.npz"),
        str(PAPER / "cache_q3.npz"): sha256(PAPER / "cache_q3.npz"),
        str(Path(__file__)): sha256(Path(__file__)),
    }
    bias_cache = IMPROVEMENT / "load_bias_cache.npz"
    if bias_cache.exists():
        input_hashes[str(bias_cache)] = sha256(bias_cache)
    manifest = {
        "command": f"python 03_代码/问题4_队友Q3优化/q4_from_q3_improvements.py --part {args.part} --start-day {args.start_day} --end-day {args.end_day}",
        "variant": "paper_q3_improvements/load_bias_2h", "seed": SEED,
        "parameters": {"issues": ISSUES, "load_bias_history_slots": 12, "load_bias_decay_slots": 36,
                       "price_forecast": "lag_7d_same_slot", "scenario_history_days": 30},
        "runtime_seconds": runtime, "python": platform.python_version(),
        "bias_cache_status": "verified_equal" if bias_cache.exists() else "absent_rebuilt_in_code",
        "inputs": input_hashes,
        "outputs": {p.name: sha256(p) for p in out.iterdir() if p.is_file() and p.name != "复现清单.json"},
    }
    (out / "复现清单.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(validation, ensure_ascii=False, indent=2))
    print(f"XLSX: {xlsx}")
    if not validation["passed"]:
        raise SystemExit(1)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--part", choices=["q42", "q43", "all"], default="all")
    ap.add_argument("--start-day", type=int, default=31)
    ap.add_argument("--end-day", type=int, default=365, help="Python右开区间；365表示运行到12月31日")
    ap.add_argument("--variant", choices=["P1", "P0"], default="P1")
    ap.add_argument("--output-dir", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()
    if not (31 <= args.start_day < args.end_day <= 365):
        raise ValueError("日期索引必须满足 31 <= start-day < end-day <= 365")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for part in (["q42", "q43"] if args.part == "all" else [args.part]):
        data, price4, daily, detail, runtime = simulate(part, args.start_day, args.end_day,
                                                        args.variant == "P1")
        write_outputs(part, daily, detail, runtime, args.output_dir, args)


if __name__ == "__main__":
    main()
