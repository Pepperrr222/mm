from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import math
import platform
import sys
import time
from dataclasses import dataclass
from datetime import datetime, time as dt_time
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import coo_matrix
from sklearn.ensemble import HistGradientBoostingRegressor


ROOT = Path(__file__).resolve().parents[1]
INPUT1 = ROOT / "01_题目与数据" / "原始附件" / "附件1.xlsx"
INPUT2 = ROOT / "01_题目与数据" / "原始附件" / "附件2.xlsx"
TEMPLATE = ROOT / "01_题目与数据" / "原始附件" / "附件5" / "result2.xlsx"
RESULTS = ROOT / "04_结果"
RESULTS.mkdir(exist_ok=True)

N = 144
H = 288
DT = 1.0 / 6.0
ETA_C = ETA_D = 0.9
E_MIN, E_MAX, E_INITIAL = 1200.0, 10800.0, 6000.0
POWER_MAX_KW = 5000.0
Q_MAX = POWER_MAX_KW * DT
EMERGENCY_MULTIPLIER = 5.0
SEED = 20260910
TAU = 1e-6
OUTPUT_START = 31  # zero-based: 2025-02-01
MODEL_NAMES = ("lag1", "mean7", "lag7", "gbdt")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def interval_label(t: int) -> str:
    a, b = t * 10, (t + 1) * 10

    def fmt(minutes: int) -> str:
        if minutes == 1440:
            return "0:00+1"
        return f"{minutes // 60}:{minutes % 60:02d}"

    return f"{fmt(a)}-{fmt(b)}"


def parse_time_minutes(value, final: bool = False) -> int:
    if isinstance(value, dt_time):
        minutes = value.hour * 60 + value.minute
    else:
        text = str(value).strip().replace("+1", "")
        hh, mm = text.split(":")[:2]
        minutes = int(hh) * 60 + int(mm)
    if final and minutes == 0:
        return 1440
    return minutes


@dataclass
class InputData:
    dates: pd.DatetimeIndex
    price: np.ndarray
    load_kw: np.ndarray
    pv_kw: np.ndarray


def read_inputs() -> InputData:
    raw1 = pd.read_excel(INPUT1, sheet_name=0, header=0)
    if raw1.shape != (N, 4):
        raise ValueError(f"附件1形状异常：{raw1.shape}，预期(144, 4)")
    expected = np.arange(1, N + 1) * 10
    observed = np.array(
        [parse_time_minutes(v, i == N - 1) for i, v in enumerate(raw1.iloc[:, 0])]
    )
    if not np.array_equal(expected, observed):
        raise ValueError("附件1时间列不是00:10至0:00+1的连续十分钟右端点")
    price = pd.to_numeric(raw1.iloc[:, 1], errors="raise").to_numpy(float)

    book = pd.ExcelFile(INPUT2)
    expected_sheets = {"小区负载", "光伏发电实际功率"}
    if set(book.sheet_names) != expected_sheets:
        raise ValueError(f"附件2工作表异常：{book.sheet_names}")
    load_df = pd.read_excel(INPUT2, sheet_name="小区负载", header=0, index_col=0)
    pv_df = pd.read_excel(INPUT2, sheet_name="光伏发电实际功率", header=0, index_col=0)
    if load_df.shape != (365, N) or pv_df.shape != (365, N):
        raise ValueError(f"附件2形状异常：load={load_df.shape}, pv={pv_df.shape}")
    dates = pd.DatetimeIndex(pd.to_datetime(load_df.index))
    if not dates.equals(pd.date_range("2025-01-01", "2025-12-31", freq="D")):
        raise ValueError("附件2日期不连续或不覆盖2025全年")
    if not dates.equals(pd.DatetimeIndex(pd.to_datetime(pv_df.index))):
        raise ValueError("附件2负载与光伏日期不一致")
    for i, col in enumerate(load_df.columns):
        if parse_time_minutes(col, i == N - 1) != expected[i]:
            raise ValueError(f"附件2第{i+2}列时间异常：{col}")
    load = load_df.apply(pd.to_numeric, errors="raise").to_numpy(float)
    pv = pv_df.apply(pd.to_numeric, errors="raise").to_numpy(float)
    if not np.isfinite(price).all() or not np.isfinite(load).all() or not np.isfinite(pv).all():
        raise ValueError("输入存在NaN或无穷值")
    if np.min(price) < 0 or np.min(load) < 0 or np.min(pv) < 0:
        raise ValueError("输入出现不合理负值")
    return InputData(dates, price, load, pv)


def date_features(date: pd.Timestamp, slots: np.ndarray) -> np.ndarray:
    doy = date.dayofyear
    dow = date.dayofweek
    return np.column_stack(
        [
            np.full_like(slots, math.sin(2 * math.pi * doy / 365.25), dtype=float),
            np.full_like(slots, math.cos(2 * math.pi * doy / 365.25), dtype=float),
            np.full_like(slots, math.sin(2 * math.pi * dow / 7), dtype=float),
            np.full_like(slots, math.cos(2 * math.pi * dow / 7), dtype=float),
            np.sin(2 * math.pi * slots / N),
            np.cos(2 * math.pi * slots / N),
        ]
    )


class CausalForecaster:
    """Daily expanding-window forecaster; all training targets are strictly before issue day."""

    def __init__(self, data: InputData, use_gbdt: bool = True):
        self.data = data
        self.use_gbdt = use_gbdt
        self.base_mae = {
            "load": {name: [] for name in MODEL_NAMES},
            "pv": {name: [] for name in MODEL_NAMES},
        }
        self.ensemble_predictions = {"load": {}, "pv": {}}
        self.residuals = {}
        self.training_cutoffs = {}

    @staticmethod
    def _simple_candidates(arr: np.ndarray, issue_day: int, offset: int, prior_pred=None):
        target = issue_day + offset
        hist = arr[:issue_day]
        if len(hist) == 0:
            raise ValueError("预测时没有历史数据")
        if offset == 0:
            lag1 = arr[issue_day - 1]
            mean_source = arr[max(0, issue_day - 7):issue_day]
        else:
            lag1 = np.asarray(prior_pred, float)
            mean_source = np.vstack([arr[max(0, issue_day - 6):issue_day], lag1])
        lag7_index = target - 7
        lag7 = arr[lag7_index] if 0 <= lag7_index < issue_day else mean_source.mean(axis=0)
        return {
            "lag1": np.asarray(lag1, float),
            "mean7": mean_source.mean(axis=0),
            "lag7": np.asarray(lag7, float),
        }

    @staticmethod
    def _gbdt(arr: np.ndarray, dates: pd.DatetimeIndex, issue_day: int, target_day: int,
              target_lag1: np.ndarray, target_lag7: np.ndarray, target_mean7: np.ndarray):
        if issue_day < 21:
            return None
        slots = np.arange(N)
        xs, ys = [], []
        for j in range(7, issue_day):
            feat = np.column_stack(
                [date_features(dates[j], slots), arr[j - 1], arr[j - 7], arr[max(0, j - 7):j].mean(axis=0)]
            )
            xs.append(feat)
            ys.append(arr[j])
        model = HistGradientBoostingRegressor(
            max_iter=35,
            max_leaf_nodes=15,
            learning_rate=0.08,
            l2_regularization=1.0,
            random_state=SEED,
        )
        model.fit(np.vstack(xs), np.concatenate(ys))
        if target_day < len(dates):
            date = dates[target_day]
        else:
            date = dates[-1] + pd.Timedelta(days=target_day - len(dates) + 1)
        x = np.column_stack(
            [date_features(date, slots), target_lag1, target_lag7, target_mean7]
        )
        return model.predict(x)

    def _weights(self, signal: str, available: list[str]) -> dict[str, float]:
        scores = {}
        for name in available:
            history = np.asarray(self.base_mae[signal][name][-28:], float)
            history = history[np.isfinite(history)]
            scores[name] = float(history.mean()) if history.size else 1.0
        inv = {name: 1.0 / max(score, 1e-6) for name, score in scores.items()}
        total = sum(inv.values())
        return {name: value / total for name, value in inv.items()}

    def _forecast_signal(self, signal: str, issue_day: int):
        arr = self.data.load_kw if signal == "load" else self.data.pv_kw
        simple0 = self._simple_candidates(arr, issue_day, 0)
        g0 = None
        if self.use_gbdt:
            g0 = self._gbdt(
                arr, self.data.dates, issue_day, issue_day,
                simple0["lag1"], simple0["lag7"], simple0["mean7"],
            )
        cand0 = dict(simple0)
        if g0 is not None:
            cand0["gbdt"] = g0
        weights = self._weights(signal, list(cand0))
        pred0 = sum(weights[name] * cand0[name] for name in cand0)
        pred0 = np.maximum(pred0, 0.0)

        simple1 = self._simple_candidates(arr, issue_day, 1, prior_pred=pred0)
        g1 = None
        if self.use_gbdt:
            g1 = self._gbdt(
                arr, self.data.dates, issue_day, issue_day + 1,
                pred0, simple1["lag7"], simple1["mean7"],
            )
        cand1 = dict(simple1)
        if g1 is not None:
            cand1["gbdt"] = g1
        pred1 = sum(weights.get(name, 0.0) * cand1[name] for name in cand1)
        if sum(weights.get(name, 0.0) for name in cand1) < 1 - 1e-12:
            denom = sum(weights.get(name, 0.0) for name in cand1)
            pred1 /= denom
        pred1 = np.maximum(pred1, 0.0)
        return pred0, pred1, cand0, weights

    def forecast(self, issue_day: int):
        load0, load1, load_candidates, load_weights = self._forecast_signal("load", issue_day)
        pv0, pv1, pv_candidates, pv_weights = self._forecast_signal("pv", issue_day)
        self.training_cutoffs[issue_day] = issue_day - 1
        return {
            "load": np.concatenate([load0, load1]),
            "pv": np.concatenate([pv0, pv1]),
            "load_day0": load0,
            "pv_day0": pv0,
            "candidates": {"load": load_candidates, "pv": pv_candidates},
            "weights": {"load": load_weights, "pv": pv_weights},
        }

    def observe(self, day: int, fc: dict):
        for signal, actual in (("load", self.data.load_kw[day]), ("pv", self.data.pv_kw[day])):
            candidates = fc["candidates"][signal]
            for name in MODEL_NAMES:
                value = float(np.mean(np.abs(actual - candidates[name]))) if name in candidates else np.nan
                self.base_mae[signal][name].append(value)
            pred = fc[f"{signal}_day0"]
            self.ensemble_predictions[signal][day] = pred.copy()
        self.residuals[day] = (
            self.data.load_kw[day] - fc["load_day0"],
            self.data.pv_kw[day] - fc["pv_day0"],
        )


def generate_scenarios(fc: dict, forecaster: CausalForecaster, issue_day: int, count: int):
    pair_starts = [j for j in sorted(forecaster.residuals) if j + 1 in forecaster.residuals and j + 1 < issue_day]
    if not pair_starts:
        load = np.repeat(fc["load"][None, :], count, axis=0)
        pv = np.repeat(fc["pv"][None, :], count, axis=0)
        return load, pv, []
    target_doy = int(forecaster.data.dates[issue_day].dayofyear)
    target_weekend = forecaster.data.dates[issue_day].dayofweek >= 5
    distances = []
    for j in pair_starts:
        date = forecaster.data.dates[j]
        season_distance = abs(int(date.dayofyear) - target_doy)
        type_penalty = 14 if (date.dayofweek >= 5) != target_weekend else 0
        distances.append(season_distance + type_penalty)
    order = np.argsort(distances)[: min(60, len(pair_starts))]
    pool = [pair_starts[i] for i in order]
    pool_dist = np.asarray([distances[i] for i in order], float)
    probability = 1.0 / (pool_dist + 7.0)
    probability /= probability.sum()
    rng = np.random.default_rng(SEED + issue_day)
    chosen = rng.choice(pool, size=count, replace=True, p=probability)
    load_s, pv_s = [], []
    for j in chosen:
        e_l = np.concatenate([forecaster.residuals[j][0], forecaster.residuals[j + 1][0]])
        e_p = np.concatenate([forecaster.residuals[j][1], forecaster.residuals[j + 1][1]])
        load_s.append(np.maximum(fc["load"] + e_l, 0.0))
        pv_s.append(np.maximum(fc["pv"] + e_p, 0.0))
    return np.asarray(load_s), np.asarray(pv_s), [int(x) for x in chosen]


def solve_stochastic_milp(price: np.ndarray, load_s_kw: np.ndarray, pv_s_kw: np.ndarray,
                          e0: float, time_limit: float, mip_gap: float):
    scenarios, horizon = load_s_kw.shape
    if horizon != H or pv_s_kw.shape != (scenarios, H):
        raise ValueError("情景形状必须为(S,288)")
    p = np.tile(price, 2)
    g0, c0, d0, e0i = 0, H, 2 * H, 3 * H
    z0 = e0i + H + 1
    h0 = z0 + H
    w0 = h0 + scenarios * H
    y0 = w0 + scenarios * H
    nv = y0 + H

    def si(base: int, s: int, t: int) -> int:
        return base + s * H + t

    objective = np.zeros(nv)
    objective[g0:g0 + H] = p
    for s in range(scenarios):
        objective[h0 + s * H:h0 + (s + 1) * H] = EMERGENCY_MULTIPLIER * p / scenarios

    lower = np.zeros(nv)
    upper = np.full(nv, np.inf)
    max_load_e = np.max(load_s_kw * DT, axis=0)
    upper[g0:g0 + H] = max_load_e + Q_MAX
    upper[c0:c0 + H] = Q_MAX
    upper[d0:d0 + H] = Q_MAX
    lower[e0i:e0i + H + 1] = E_MIN
    upper[e0i:e0i + H + 1] = E_MAX
    upper[z0:z0 + H] = 1.0
    mh = 1.05 * float(np.max(load_s_kw * DT))
    upper[h0:h0 + scenarios * H] = mh
    upper[y0:y0 + H] = 1.0
    integrality = np.zeros(nv, dtype=np.uint8)
    integrality[z0:z0 + H] = 1
    integrality[y0:y0 + H] = 1

    rows, cols, vals, lo, hi = [], [], [], [], []

    def add(entries, low=-np.inf, high=np.inf):
        r = len(lo)
        for col, value in entries:
            rows.append(r)
            cols.append(col)
            vals.append(value)
        lo.append(low)
        hi.append(high)

    for t in range(H):
        add([(e0i + t + 1, 1), (e0i + t, -1), (c0 + t, -ETA_C), (d0 + t, 1 / ETA_D)], 0, 0)
        add([(c0 + t, 1), (z0 + t, -Q_MAX)], high=0)
        add([(d0 + t, 1), (z0 + t, Q_MAX)], high=Q_MAX)
        add([(c0 + t, 1), (y0 + t, Q_MAX)], high=Q_MAX)
    add([(e0i, 1)], e0, e0)
    add([(e0i + H, 1)], e0, e0)
    for s in range(scenarios):
        net = (load_s_kw[s] - pv_s_kw[s]) * DT
        for t in range(H):
            add(
                [(g0 + t, 1), (c0 + t, -1), (d0 + t, 1),
                 (si(h0, s, t), 1), (si(w0, s, t), -1)],
                net[t], net[t],
            )
            add([(si(h0, s, t), 1), (y0 + t, -mh)], high=0)
    matrix = coo_matrix((vals, (rows, cols)), shape=(len(lo), nv)).tocsr()
    started = time.perf_counter()
    result = milp(
        objective,
        integrality=integrality,
        bounds=Bounds(lower, upper),
        constraints=LinearConstraint(matrix, np.asarray(lo), np.asarray(hi)),
        options={"time_limit": time_limit, "mip_rel_gap": mip_gap, "presolve": True},
    )
    elapsed = time.perf_counter() - started
    if not result.success or result.x is None:
        raise RuntimeError(f"HiGHS失败 status={result.status}: {result.message}")
    x = result.x
    plan = {
        "grid": x[g0:g0 + H],
        "charge": x[c0:c0 + H],
        "discharge": x[d0:d0 + H],
        "E": x[e0i:e0i + H + 1],
        "z": x[z0:z0 + H],
        "h": x[h0:h0 + scenarios * H].reshape(scenarios, H),
        "w": x[w0:w0 + scenarios * H].reshape(scenarios, H),
        "objective": float(result.fun),
        "elapsed_s": elapsed,
        "mip_gap": float(getattr(result, "mip_gap", np.nan)),
        "scenario_count": scenarios,
        "mh": mh,
    }
    return plan


def solve_with_fallback(price, load_s, pv_s, e0, time_limit, mip_gap):
    attempts = [(load_s, pv_s, "main")]
    if len(load_s) > 10:
        attempts.append((load_s[:10], pv_s[:10], "10_scenario"))
    conservative_load = np.quantile(load_s, 0.8, axis=0, method="higher")[None, :]
    conservative_pv = np.quantile(pv_s, 0.2, axis=0, method="lower")[None, :]
    attempts.append((conservative_load, conservative_pv, "deterministic_q80"))
    errors = []
    for a_load, a_pv, tag in attempts:
        try:
            plan = solve_stochastic_milp(price, a_load, a_pv, e0, time_limit, mip_gap)
            plan["fallback"] = tag
            plan["validation_load_s"] = a_load
            plan["validation_pv_s"] = a_pv
            return plan
        except RuntimeError as exc:
            errors.append(f"{tag}: {exc}")
    raise RuntimeError("全部求解尝试失败；" + " | ".join(errors))


def execute_interval(e, g, c_plan, d_plan, load_e, pv_e):
    """Causal real-time safety projection using only the current observed interval."""
    d_cap = max(0.0, (e - E_MIN) * ETA_D)
    current_deficit = max(0.0, load_e - g - pv_e)
    d_exec = min(max(d_plan, 0.0), d_cap, Q_MAX, current_deficit)
    current_surplus = max(0.0, g + pv_e - load_e)
    c_cap = max(0.0, (E_MAX - e) / ETA_C)
    c_exec = min(max(c_plan, 0.0), current_surplus, Q_MAX, c_cap)
    emergency = max(0.0, load_e + c_exec - g - pv_e - d_exec)
    surplus = max(0.0, g + pv_e + d_exec - load_e - c_exec)
    e_next = e + ETA_C * c_exec - d_exec / ETA_D
    return c_exec, d_exec, emergency, surplus, e_next


def run_execution_unit_tests():
    # Partial charge: only the observed 400 kWh surplus may enter the battery.
    c, d, h, w, e1 = execute_interval(6000, 1000, Q_MAX, 0, 600, 0)
    assert abs(c - 400) < 1e-9 and d == 0 and h == 0 and w == 0 and abs(e1 - 6360) < 1e-9
    # Planned charging is fully cancelled when the current interval has a shortage.
    c, d, h, w, e1 = execute_interval(6000, 100, Q_MAX, 0, 600, 0)
    assert c == 0 and d == 0 and abs(h - 500) < 1e-9 and w == 0 and abs(e1 - 6000) < 1e-9
    # Planned discharge is clipped to the observed deficit, so it is not wasted.
    c, d, h, w, e1 = execute_interval(6000, 100, 0, Q_MAX, 500, 0)
    assert c == 0 and abs(d - 400) < 1e-9 and h == 0 and w == 0
    assert abs(e1 - (6000 - 400 / ETA_D)) < 1e-9


def replay_day(date, price, load_kw, pv_kw, plan, e_start):
    rows = []
    e = float(e_start)
    for t in range(N):
        g = max(float(plan["grid"][t]), 0.0)
        c_plan = max(float(plan["charge"][t]), 0.0)
        d_plan = max(float(plan["discharge"][t]), 0.0)
        load_e = float(load_kw[t] * DT)
        pv_e = float(pv_kw[t] * DT)
        c_exec, d_exec, emergency, surplus, e_next = execute_interval(
            e, g, c_plan, d_plan, load_e, pv_e
        )
        rows.append(
            {
                "date": pd.Timestamp(date), "t": t, "interval": interval_label(t),
                "price": float(price[t]), "load_kw": float(load_kw[t]), "pv_kw": float(pv_kw[t]),
                "grid_kwh": g, "charge_plan_kwh": c_plan, "discharge_plan_kwh": d_plan,
                "charge_exec_kwh": c_exec, "discharge_exec_kwh": d_exec,
                "emergency_kwh": emergency, "surplus_kwh": surplus,
                "E_start_kwh": e, "E_end_kwh": e_next,
                "plan_cost_yuan": float(price[t] * g),
                "emergency_cost_yuan": float(EMERGENCY_MULTIPLIER * price[t] * emergency),
            }
        )
        e = e_next
    return pd.DataFrame(rows), e


def validate_day(real: pd.DataFrame, plan: dict, e_start: float,
                 load_s_kw: np.ndarray, pv_s_kw: np.ndarray):
    balance = (
        real.grid_kwh + real.pv_kw * DT + real.discharge_exec_kwh + real.emergency_kwh
        - real.load_kw * DT - real.charge_exec_kwh - real.surplus_kwh
    )
    state = (
        real.E_end_kwh - real.E_start_kwh - ETA_C * real.charge_exec_kwh
        + real.discharge_exec_kwh / ETA_D
    )
    plan_state = (
        plan["E"][1:] - plan["E"][:-1] - ETA_C * plan["charge"] + plan["discharge"] / ETA_D
    )
    scenario_balance = (
        plan["grid"][None, :] + pv_s_kw * DT + plan["discharge"][None, :] + plan["h"]
        - load_s_kw * DT - plan["charge"][None, :] - plan["w"]
    )
    metrics = {
        "max_real_balance_residual_kwh": float(np.max(np.abs(balance))),
        "max_real_state_residual_kwh": float(np.max(np.abs(state))),
        "max_plan_state_residual_kwh": float(np.max(np.abs(plan_state))),
        "max_scenario_balance_residual_kwh": float(np.max(np.abs(scenario_balance))),
        "real_E_min_kwh": float(min(real.E_start_kwh.min(), real.E_end_kwh.min())),
        "real_E_max_kwh": float(max(real.E_start_kwh.max(), real.E_end_kwh.max())),
        "plan_E_min_kwh": float(np.min(plan["E"])),
        "plan_E_max_kwh": float(np.max(plan["E"])),
        "plan_terminal_gap_kwh": float(plan["E"][-1] - e_start),
        "simultaneous_plan_periods": int(np.sum((plan["charge"] > TAU) & (plan["discharge"] > TAU))),
        "emergency_charge_exec_periods": int(np.sum((real.emergency_kwh > TAU) & (real.charge_exec_kwh > TAU))),
        "emergency_charge_scenario_periods": int(np.sum((plan["h"] > TAU) & (plan["charge"][None, :] > TAU))),
        "emergency_surplus_scenario_periods": int(np.sum((plan["h"] > TAU) & (plan["w"] > TAU))),
        "max_charge_exec_kwh": float(real.charge_exec_kwh.max()),
        "max_discharge_exec_kwh": float(real.discharge_exec_kwh.max()),
    }
    if metrics["max_real_balance_residual_kwh"] > 1e-5 or metrics["max_scenario_balance_residual_kwh"] > 1e-5:
        raise RuntimeError(f"实际能量平衡失败：{metrics}")
    if metrics["max_real_state_residual_kwh"] > 1e-5 or metrics["max_plan_state_residual_kwh"] > 1e-5:
        raise RuntimeError(f"储能递推失败：{metrics}")
    if metrics["real_E_min_kwh"] < E_MIN - 1e-5 or metrics["real_E_max_kwh"] > E_MAX + 1e-5:
        raise RuntimeError(f"实际储能越界：{metrics}")
    if metrics["plan_E_min_kwh"] < E_MIN - 1e-5 or metrics["plan_E_max_kwh"] > E_MAX + 1e-5:
        raise RuntimeError(f"计划储能越界：{metrics}")
    if (metrics["simultaneous_plan_periods"] or metrics["emergency_charge_exec_periods"]
            or metrics["emergency_charge_scenario_periods"] or metrics["emergency_surplus_scenario_periods"]):
        raise RuntimeError(f"互斥约束失败：{metrics}")
    return metrics


def emergency_episodes(detail: pd.DataFrame) -> pd.DataFrame:
    records = []
    for date, part in detail.groupby("date", sort=True):
        positive = part.emergency_kwh.to_numpy() > TAU
        starts = np.flatnonzero(positive & ~np.r_[False, positive[:-1]])
        ends = np.flatnonzero(positive & ~np.r_[positive[1:], False])
        if not len(starts):
            records.append({"date": date, "interval": "无", "emergency_kwh": 0.0})
        for a, b in zip(starts, ends):
            left = interval_label(int(a)).split("-")[0]
            right = interval_label(int(b)).split("-")[1]
            records.append(
                {"date": date, "interval": f"{left}-{right}",
                 "emergency_kwh": float(part.iloc[a:b + 1].emergency_kwh.sum())}
            )
    return pd.DataFrame(records)


def write_full_workbook(detail: pd.DataFrame, daily: pd.DataFrame):
    output = RESULTS / "result2_问题2_结果.xlsx"
    wb = openpyxl.load_workbook(TEMPLATE)
    plan_ws = wb["计划购电量"]
    headers = [interval_label(t) for t in range(N)]
    for t, value in enumerate(headers, 2):
        plan_ws.cell(1, t).value = value
    for r, (date, part) in enumerate(detail.groupby("date", sort=True), 2):
        row_metric = daily.loc[daily.date == date].iloc[0]
        plan_ws.cell(r, 1).value = pd.Timestamp(date).to_pydatetime()
        for t, value in enumerate(part.grid_kwh, 2):
            plan_ws.cell(r, t).value = float(value)
        plan_ws.cell(r, 146).value = float(part.grid_kwh.sum())
        plan_ws.cell(r, 147).value = float(row_metric.plan_cost_yuan)

    charge_ws = wb["充放电量"]
    if charge_ws.max_row > 1:
        charge_ws.delete_rows(2, charge_ws.max_row - 1)
    row = 2
    for date, part in detail.groupby("date", sort=True):
        for block in range(6):
            sub = part.iloc[block * 24:(block + 1) * 24]
            charge_ws.cell(row, 1).value = pd.Timestamp(date).to_pydatetime() if block == 0 else None
            charge_ws.cell(row, 2).value = f"{block * 4}:00-{(block + 1) * 4}:00"
            charge_ws.cell(row, 3).value = float(sub.charge_exec_kwh.sum())
            charge_ws.cell(row, 4).value = float(sub.discharge_exec_kwh.sum())
            if block == 0:
                charge_ws.cell(row, 5).value = "0:00"
                charge_ws.cell(row, 6).value = float(part.iloc[0].E_start_kwh)
            elif block == 1:
                charge_ws.cell(row, 5).value = "24:00"
                charge_ws.cell(row, 6).value = float(part.iloc[-1].E_end_kwh)
            row += 1

    emg_ws = wb["紧急购电量"]
    if emg_ws.max_row > 1:
        emg_ws.delete_rows(2, emg_ws.max_row - 1)
    episodes = emergency_episodes(detail)
    row = 2
    for date, part in episodes.groupby("date", sort=True):
        for i, rec in enumerate(part.itertuples(index=False)):
            emg_ws.cell(row, 1).value = pd.Timestamp(date).to_pydatetime() if i == 0 else None
            emg_ws.cell(row, 2).value = rec.interval
            emg_ws.cell(row, 3).value = float(rec.emergency_kwh)
            row += 1
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(output)
    return output


def run(args):
    run_execution_unit_tests()
    data = read_inputs()
    forecaster = CausalForecaster(data, use_gbdt=not args.no_gbdt)
    end_day = min(len(data.dates), OUTPUT_START + args.days)
    e_real = E_INITIAL
    detail_parts, daily_rows, forecast_rows, validation_rows = [], [], [], []

    # January is a causal calibration period; storage remains idle, so E=6000 on Feb 1.
    for day in range(1, end_day):
        fc = forecaster.forecast(day)
        if day < OUTPUT_START:
            forecaster.observe(day, fc)
            continue
        load_s, pv_s, sampled = generate_scenarios(fc, forecaster, day, args.scenarios)
        plan = solve_with_fallback(
            data.price, load_s, pv_s, e_real, args.solver_time_limit, args.mip_gap
        )
        real, e_next = replay_day(
            data.dates[day], data.price, data.load_kw[day], data.pv_kw[day], plan, e_real
        )
        checks = validate_day(
            real, plan, e_real, plan["validation_load_s"], plan["validation_pv_s"]
        )
        detail_parts.append(real)
        validation_rows.append({"date": data.dates[day], **checks})
        daily_rows.append(
            {
                "date": data.dates[day],
                "E_start_kwh": e_real,
                "E_end_kwh": e_next,
                "plan_grid_kwh": float(real.grid_kwh.sum()),
                "plan_cost_yuan": float(real.plan_cost_yuan.sum()),
                "emergency_kwh": float(real.emergency_kwh.sum()),
                "emergency_cost_yuan": float(real.emergency_cost_yuan.sum()),
                "total_cost_yuan": float(real.plan_cost_yuan.sum() + real.emergency_cost_yuan.sum()),
                "surplus_kwh": float(real.surplus_kwh.sum()),
                "charge_exec_kwh": float(real.charge_exec_kwh.sum()),
                "discharge_exec_kwh": float(real.discharge_exec_kwh.sum()),
                "solver_objective_48h_yuan": plan["objective"],
                "solver_elapsed_s": plan["elapsed_s"],
                "solver_mip_gap": plan["mip_gap"],
                "scenario_count_used": plan["scenario_count"],
                "fallback": plan["fallback"],
                "sampled_residual_days": json.dumps(sampled, ensure_ascii=False),
                "sampled_residual_dates": json.dumps(
                    [str(data.dates[j].date()) for j in sampled], ensure_ascii=False
                ),
                "training_cutoff": data.dates[forecaster.training_cutoffs[day]],
            }
        )
        for signal in ("load", "pv"):
            actual = data.load_kw[day] if signal == "load" else data.pv_kw[day]
            pred = fc[f"{signal}_day0"]
            forecast_rows.append(
                {
                    "date": data.dates[day], "signal": signal,
                    "mae_kw": float(np.mean(np.abs(actual - pred))),
                    "rmse_kw": float(np.sqrt(np.mean((actual - pred) ** 2))),
                    "nmae": float(np.mean(np.abs(actual - pred)) / max(np.mean(actual), 1e-9)),
                    "weights": json.dumps(fc["weights"][signal], ensure_ascii=False),
                    "training_cutoff": data.dates[day - 1],
                }
            )
        e_real = e_next
        forecaster.observe(day, fc)
        print(
            f"[{data.dates[day].date()}] cost={daily_rows[-1]['total_cost_yuan']:.2f}, "
            f"emg={daily_rows[-1]['emergency_kwh']:.2f} kWh, E={e_next:.2f}, "
            f"solver={plan['elapsed_s']:.2f}s, mode={plan['fallback']}",
            flush=True,
        )

    if not detail_parts:
        raise RuntimeError("没有进入正式输出日期，请增加--days")
    detail = pd.concat(detail_parts, ignore_index=True)
    daily = pd.DataFrame(daily_rows)
    forecast_metrics = pd.DataFrame(forecast_rows)
    validation = pd.DataFrame(validation_rows)
    prefix = "问题2" if args.days >= 334 else "问题2_P1"
    outputs = []
    for frame, name in (
        (detail, f"{prefix}_逐时段结果.csv"),
        (daily, f"{prefix}_每日指标.csv"),
        (forecast_metrics, f"{prefix}_预测指标.csv"),
        (validation, f"{prefix}_验证指标.csv"),
        (emergency_episodes(detail), f"{prefix}_紧急购电区间.csv"),
    ):
        path = RESULTS / name
        frame.to_csv(path, index=False, encoding="utf-8-sig")
        outputs.append(path)
    workbook = None
    if args.days >= 334:
        workbook = write_full_workbook(detail, daily)
        outputs.append(workbook)

    summary = {
        "days": int(len(daily)),
        "date_start": str(daily.date.min().date()),
        "date_end": str(daily.date.max().date()),
        "total_cost_yuan": float(daily.total_cost_yuan.sum()),
        "plan_cost_yuan": float(daily.plan_cost_yuan.sum()),
        "emergency_cost_yuan": float(daily.emergency_cost_yuan.sum()),
        "emergency_kwh": float(daily.emergency_kwh.sum()),
        "surplus_kwh": float(daily.surplus_kwh.sum()),
        "E_min_kwh": float(min(detail.E_start_kwh.min(), detail.E_end_kwh.min())),
        "E_max_kwh": float(max(detail.E_start_kwh.max(), detail.E_end_kwh.max())),
        "max_balance_residual_kwh": float(validation.max_real_balance_residual_kwh.max()),
        "max_state_residual_kwh": float(validation.max_real_state_residual_kwh.max()),
        "max_cross_day_state_gap_kwh": float(np.max(np.abs(daily.E_start_kwh.iloc[1:].to_numpy() - daily.E_end_kwh.iloc[:-1].to_numpy()))) if len(daily) > 1 else 0.0,
        "fallback_counts": daily.fallback.value_counts().to_dict(),
    }
    manifest = {
        "generated_at": datetime.now().isoformat(),
        "python": sys.version,
        "platform": platform.platform(),
        "command": "python " + " ".join(sys.argv),
        "seed": SEED,
        "inputs": {str(p.relative_to(ROOT)): sha256(p) for p in (INPUT1, INPUT2, TEMPLATE)},
        "code_sha256": sha256(Path(__file__)),
        "dependencies": {
            name: importlib.metadata.version(name)
            for name in ["numpy", "pandas", "scipy", "scikit-learn", "openpyxl"]
        },
        "outputs": {str(p.relative_to(ROOT)): sha256(p) for p in outputs},
        "parameters": {
            "N": N, "H": H, "dt_h": DT, "eta_c": ETA_C, "eta_d": ETA_D,
            "E_min_kwh": E_MIN, "E_max_kwh": E_MAX, "E_initial_kwh": E_INITIAL,
            "power_max_kw": POWER_MAX_KW, "emergency_multiplier": EMERGENCY_MULTIPLIER,
            "scenarios": args.scenarios, "mip_gap": args.mip_gap,
            "solver_time_limit_s": args.solver_time_limit, "gbdt": not args.no_gbdt,
        },
        "summary": summary,
    }
    manifest_path = RESULTS / f"{prefix}_复现清单.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return summary


def parse_args():
    parser = argparse.ArgumentParser(description="C题问题2：因果预测、情景随机MILP与真实回放")
    parser.add_argument("--days", type=int, default=334, help="从2025-02-01起运行天数")
    parser.add_argument("--scenarios", type=int, default=30, help="历史联合残差情景数")
    parser.add_argument("--solver-time-limit", type=float, default=30.0)
    parser.add_argument("--mip-gap", type=float, default=1e-6)
    parser.add_argument("--no-gbdt", action="store_true", help="仅用于诊断的轻量预测模式")
    return parser.parse_args()


if __name__ == "__main__":
    run(parse_args())
