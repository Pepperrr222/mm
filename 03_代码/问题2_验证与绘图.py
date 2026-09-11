#!/usr/bin/env python3
"""问题2：对全年结果做独立复算、模板核验并生成论文图。"""

from __future__ import annotations

import hashlib
import json
import shutil
import sys
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RESULT = ROOT / "04_结果"
FIGURE = ROOT / "05_图表"
PREVIEW = FIGURE / "previews"
STYLE_SRC = Path(r"C:\Users\Lenovo\.codex\skills\math-modeling\references\roles\编程手\scripts\plot_style.py")
EXPORT_DIR = Path(r"C:\Users\Lenovo\.codex\skills\math-modeling\tools\figure\scripts")
if str(EXPORT_DIR) not in sys.path:
    sys.path.insert(0, str(EXPORT_DIR))
from export_figure import export_figure


PALETTE = {
    "blue": "#0072B2", "orange": "#E69F00", "green": "#009E73",
    "red": "#D55E00", "purple": "#CC79A7", "sky": "#56B4E9",
    "gray": "#6B7280", "dark": "#222222",
}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def style() -> None:
    plt.rcParams.update({
        "figure.figsize": (7.2, 4.4), "figure.dpi": 120,
        "font.family": "DejaVu Sans", "font.size": 8,
        "axes.titlesize": 9, "axes.labelsize": 8,
        "xtick.labelsize": 7, "ytick.labelsize": 7,
        "legend.fontsize": 7, "axes.spines.top": False,
        "axes.spines.right": False, "axes.linewidth": 0.7,
        "grid.alpha": 0.22, "grid.linewidth": 0.5,
        "svg.fonttype": "none", "pdf.fonttype": 42,
    })


def save(fig: plt.Figure, stem: str, size=(7.2, 4.4)) -> list[str]:
    paths = export_figure(
        fig, str(FIGURE / stem), formats=("svg", "png"), dpi=300,
        size_inches=size, grayscale_preview=True, tight=True,
    )
    plt.close(fig)
    PREVIEW.mkdir(parents=True, exist_ok=True)
    gray = FIGURE / f"{stem}_grayscale.png"
    if gray.exists():
        target = PREVIEW / gray.name
        if target.exists():
            target.unlink()
        shutil.move(str(gray), str(target))
        paths = [str(target) if Path(p) == gray else p for p in paths]
    return paths


def validate(detail: pd.DataFrame, daily: pd.DataFrame, forecast: pd.DataFrame,
             check: pd.DataFrame, workbook: Path) -> dict:
    dt = 1 / 6
    balance = (detail.grid_kwh + detail.emergency_kwh + detail.discharge_exec_kwh
               + detail.pv_kw * dt - detail.charge_exec_kwh
               - detail.load_kw * dt - detail.surplus_kwh)
    state = detail.E_end_kwh - (
        detail.E_start_kwh + 0.9 * detail.charge_exec_kwh
        - detail.discharge_exec_kwh / 0.9
    )
    continuity = detail.groupby("date").agg(start=("E_start_kwh", "first"), end=("E_end_kwh", "last"))
    cross = continuity.start.iloc[1:].to_numpy() - continuity.end.iloc[:-1].to_numpy()
    price_grid_cost = float((detail.price * detail.grid_kwh).sum())
    emergency_cost = float((5 * detail.price * detail.emergency_kwh).sum())

    wb = load_workbook(workbook, read_only=False, data_only=False)
    sheet_info = {ws.title: [ws.max_row, ws.max_column] for ws in wb.worksheets}
    formula_errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#"):
                    formula_errors.append(f"{ws.title}!{cell.coordinate}:{cell.value}")
    ws = wb["计划购电量"]
    dates = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
    required_dates = {pd.Timestamp(x).date() for x in ("2025-02-01", "2025-05-01", "2025-08-01", "2025-11-01")}
    workbook_dates = {pd.Timestamp(x).date() for x in dates if x is not None}

    metrics = {
        "rows_detail": int(len(detail)), "days": int(detail.date.nunique()),
        "date_start": str(detail.date.min().date()), "date_end": str(detail.date.max().date()),
        "rows_per_day_min": int(detail.groupby("date").size().min()),
        "rows_per_day_max": int(detail.groupby("date").size().max()),
        "plan_cost_recomputed_yuan": price_grid_cost,
        "plan_cost_reported_yuan": float(daily.plan_cost_yuan.sum()),
        "emergency_cost_recomputed_yuan": emergency_cost,
        "emergency_cost_reported_yuan": float(daily.emergency_cost_yuan.sum()),
        "total_cost_yuan": float(daily.total_cost_yuan.sum()),
        "max_balance_residual_kwh": float(np.abs(balance).max()),
        "max_state_residual_kwh": float(np.abs(state).max()),
        "max_cross_day_state_gap_kwh": float(np.abs(cross).max()),
        "energy_min_kwh": float(detail[["E_start_kwh", "E_end_kwh"]].min().min()),
        "energy_max_kwh": float(detail[["E_start_kwh", "E_end_kwh"]].max().max()),
        "simultaneous_plan_periods": int(((detail.charge_plan_kwh > 1e-7) & (detail.discharge_plan_kwh > 1e-7)).sum()),
        "emergency_charge_periods": int(((detail.emergency_kwh > 1e-7) & (detail.charge_exec_kwh > 1e-7)).sum()),
        "forecast_rows": int(len(forecast)), "validation_days": int(len(check)),
        "max_plan_state_residual_kwh": float(check.max_plan_state_residual_kwh.max()),
        "max_scenario_balance_residual_kwh": float(check.max_scenario_balance_residual_kwh.max()),
        "validation_simultaneous_plan_periods": int(check.simultaneous_plan_periods.sum()),
        "validation_emergency_charge_scenario_periods": int(check.emergency_charge_scenario_periods.sum()),
        "validation_emergency_surplus_scenario_periods": int(check.emergency_surplus_scenario_periods.sum()),
        "fallback_counts": {str(k): int(v) for k, v in daily.fallback.value_counts().items()},
        "workbook_sheets": sheet_info,
        "workbook_formula_error_count": len(formula_errors),
        "workbook_required_dates_present": required_dates.issubset(workbook_dates),
        "workbook_first_interval": ws.cell(1, 2).value,
        "workbook_last_interval": ws.cell(1, 145).value,
    }
    assertions = {
        "complete_334_days": metrics["days"] == 334 and metrics["rows_detail"] == 334 * 144,
        "date_range_ok": metrics["date_start"] == "2025-02-01" and metrics["date_end"] == "2025-12-31",
        "cost_recompute_ok": abs(price_grid_cost - metrics["plan_cost_reported_yuan"]) < 1e-5 and abs(emergency_cost - metrics["emergency_cost_reported_yuan"]) < 1e-5,
        "balance_ok": metrics["max_balance_residual_kwh"] < 1e-8,
        "state_ok": metrics["max_state_residual_kwh"] < 1e-8 and metrics["max_cross_day_state_gap_kwh"] < 1e-8,
        "energy_bounds_ok": metrics["energy_min_kwh"] >= 1200 - 1e-7 and metrics["energy_max_kwh"] <= 10800 + 1e-7,
        "logic_ok": metrics["simultaneous_plan_periods"] == 0 and metrics["emergency_charge_periods"] == 0,
        "planned_and_scenario_residuals_ok": metrics["max_plan_state_residual_kwh"] < 1e-5 and metrics["max_scenario_balance_residual_kwh"] < 1e-5,
        "scenario_logic_ok": metrics["validation_simultaneous_plan_periods"] == 0 and metrics["validation_emergency_charge_scenario_periods"] == 0 and metrics["validation_emergency_surplus_scenario_periods"] == 0,
        "workbook_ok": sheet_info == {"计划购电量": [335, 147], "充放电量": [2005, 6], "紧急购电量": [3592, 3]} and len(formula_errors) == 0 and required_dates.issubset(workbook_dates),
    }
    return {"status": "PASS" if all(assertions.values()) else "FAIL", "assertions": assertions, "metrics": metrics, "formula_errors": formula_errors[:20]}


def create_figures(detail: pd.DataFrame, daily: pd.DataFrame, forecast: pd.DataFrame) -> list[str]:
    detail = detail.copy()
    detail["date"] = pd.to_datetime(detail.date)
    daily = daily.copy()
    daily["date"] = pd.to_datetime(daily.date)
    forecast = forecast.copy()
    forecast["date"] = pd.to_datetime(forecast.date)
    detail["month"] = detail.date.dt.month
    daily["month"] = daily.date.dt.month
    outputs: list[str] = []

    monthly_energy = detail.groupby("month").agg(load_kw=("load_kw", "sum"), pv_kw=("pv_kw", "sum")) / 6000
    fig, ax = plt.subplots()
    ax.plot(monthly_energy.index, monthly_energy.load_kw, color=PALETTE["blue"], marker="o", label="Load")
    ax.plot(monthly_energy.index, monthly_energy.pv_kw, color=PALETTE["orange"], marker="s", label="PV")
    ax.set(xlabel="Month", ylabel="Energy (MWh)", title="Monthly load and photovoltaic energy")
    ax.set_xticks(monthly_energy.index); ax.grid(axis="y"); ax.legend()
    outputs += save(fig, "raw_q2_monthly_energy")

    heat = detail.pivot(index="date", columns="t", values="load_kw") - detail.pivot(index="date", columns="t", values="pv_kw")
    fig, ax = plt.subplots()
    im = ax.imshow(heat.to_numpy() / 1000, aspect="auto", origin="lower", cmap="RdBu_r")
    ax.set(xlabel="Time of day", ylabel="Days since 2025-02-01", title="Actual net-load calendar")
    ax.set_xticks([0, 36, 72, 108, 143], ["00", "06", "12", "18", "24"])
    fig.colorbar(im, ax=ax, label="Net load (MW)")
    outputs += save(fig, "raw_q2_netload_heatmap")

    fig, ax = plt.subplots()
    data = [forecast.loc[forecast.signal == s, "nmae"].to_numpy() * 100 for s in ("load", "pv")]
    bp = ax.boxplot(data, tick_labels=["Load", "PV"], patch_artist=True, showfliers=False)
    for box, c in zip(bp["boxes"], (PALETTE["blue"], PALETTE["orange"])): box.set_facecolor(c); box.set_alpha(.65)
    rng = np.random.default_rng(20260911)
    for i, y in enumerate(data, 1): ax.scatter(rng.normal(i, .045, len(y)), y, s=5, alpha=.22, color=PALETTE["dark"])
    ax.set(ylabel="Daily NMAE (%)", title="Causal one-day-ahead forecast error"); ax.grid(axis="y")
    outputs += save(fig, "raw_q2_forecast_error_distribution")

    counts = daily.groupby(["month", "fallback"]).size().unstack(fill_value=0)
    order = [c for c in ("main", "10_scenario", "deterministic_q80") if c in counts]
    short_mode = {"main": "30-scenario", "10_scenario": "10-scenario", "deterministic_q80": "Q80 fallback"}
    fig, ax = plt.subplots(); bottom = np.zeros(len(counts))
    for c, color, hatch in zip(order, (PALETTE["blue"], PALETTE["orange"], PALETTE["red"]), ("//", "..", "xx")):
        vals = counts[c].to_numpy()
        ax.bar(counts.index, vals, bottom=bottom, label=short_mode[c], color=color,
               hatch=hatch, edgecolor="white", linewidth=.5)
        bottom += vals
    ax.set(xlabel="Month", ylabel="Days"); ax.set_xticks(counts.index); ax.grid(axis="y")
    handles, labels = ax.get_legend_handles_labels()
    fig.subplots_adjust(top=.78)
    fig.suptitle("Optimization mode used by month", y=.98)
    fig.legend(handles, labels, loc="upper center", bbox_to_anchor=(.5, .91), ncol=3, frameon=False)
    outputs += save(fig, "process_q2_solver_modes")

    fig, ax = plt.subplots()
    ax.fill_between(daily.date, 1200, 10800, color=PALETTE["sky"], alpha=.10, label="Feasible range")
    ax.plot(daily.date, daily.E_start_kwh, color=PALETTE["gray"], lw=.9, ls="--", label="Start")
    ax.plot(daily.date, daily.E_end_kwh, color=PALETTE["green"], lw=1.0, label="End")
    ax.set(ylabel="Stored energy (kWh)", title="Daily storage state and inherited continuity"); ax.grid(axis="y")
    ax.legend(ncol=3, loc="lower left")
    outputs += save(fig, "process_q2_storage_state")

    fig, ax = plt.subplots()
    for i, mode in enumerate(order, 1):
        y = daily.loc[daily.fallback == mode, "solver_elapsed_s"].to_numpy()
        x = np.full(len(y), i, float) + np.random.default_rng(20260911 + i).normal(0, .055, len(y))
        ax.scatter(x, y, s=10, alpha=.5, color=(PALETTE["blue"], PALETTE["orange"], PALETTE["red"])[i-1])
        if len(y) >= 3: ax.hlines(np.median(y), i-.22, i+.22, color=PALETTE["dark"], lw=1.4)
    ax.set_yscale("log"); ax.set_xticks(range(1, len(order)+1), [short_mode[x] for x in order])
    ax.set(ylabel="Solver time (s, log scale)", title="Daily solver time with all observations"); ax.grid(axis="y")
    outputs += save(fig, "process_q2_solver_time")

    month_cost = daily.groupby("month")[["plan_cost_yuan", "emergency_cost_yuan"]].sum() / 1e6
    fig, ax = plt.subplots(); ax.bar(month_cost.index, month_cost.plan_cost_yuan, color=PALETTE["blue"], label="Planned")
    ax.bar(month_cost.index, month_cost.emergency_cost_yuan, bottom=month_cost.plan_cost_yuan, color=PALETTE["red"], label="Emergency")
    ax.set(xlabel="Month", ylabel="Cost (million CNY)", title="Realized monthly electricity cost"); ax.set_xticks(month_cost.index); ax.legend(); ax.grid(axis="y")
    outputs += save(fig, "result_q2_monthly_cost")

    month_emergency = daily.groupby("month").emergency_kwh.sum() / 1000
    fig, ax = plt.subplots(); ax.bar(month_emergency.index, month_emergency, color=PALETTE["red"])
    ax.set(xlabel="Month", ylabel="Emergency energy (MWh)", title="Monthly emergency procurement"); ax.set_xticks(month_emergency.index); ax.grid(axis="y")
    outputs += save(fig, "result_q2_emergency_energy")

    dates = pd.to_datetime(["2025-02-01", "2025-05-01", "2025-08-01", "2025-11-01"])
    fig, axes = plt.subplots(2, 2, figsize=(7.2, 5.2), sharex=True)
    for label, ax, day in zip("abcd", axes.ravel(), dates):
        z = detail.loc[detail.date == day]
        ax.plot(z.t / 6, z.grid_kwh * 6 / 1000, color=PALETTE["blue"], label="Planned grid")
        ax.fill_between(z.t / 6, 0, z.emergency_kwh * 6 / 1000, color=PALETTE["red"], alpha=.55, label="Emergency")
        ax.text(-.11, 1.02, label, transform=ax.transAxes, fontweight="bold")
        ax.set_title(str(day.date())); ax.grid(axis="y")
    for ax in axes[-1]: ax.set_xlabel("Hour")
    for ax in axes[:, 0]: ax.set_ylabel("Power-equivalent (MW)")
    handles, labels = axes[0, 0].get_legend_handles_labels()
    fig.subplots_adjust(top=.82, hspace=.34, wspace=.20)
    fig.suptitle("Specified-day realized procurement", y=.985)
    fig.legend(handles, labels, loc="upper center", bbox_to_anchor=(.5, .93), ncol=2, frameon=False)
    outputs += save(fig, "result_q2_specified_dates", size=(7.2, 5.2))
    return outputs


def main() -> int:
    (ROOT / "utils").mkdir(exist_ok=True)
    if STYLE_SRC.exists():
        shutil.copy2(STYLE_SRC, ROOT / "utils" / "plot_style.py")
    FIGURE.mkdir(exist_ok=True)
    style()
    detail = pd.read_csv(RESULT / "问题2_逐时段结果.csv", parse_dates=["date"])
    daily = pd.read_csv(RESULT / "问题2_每日指标.csv", parse_dates=["date", "training_cutoff"])
    forecast = pd.read_csv(RESULT / "问题2_预测指标.csv", parse_dates=["date", "training_cutoff"])
    check = pd.read_csv(RESULT / "问题2_验证指标.csv", parse_dates=["date"])
    report = validate(detail, daily, forecast, check, RESULT / "result2_问题2_结果.xlsx")
    (RESULT / "问题2_总体验证.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    specified = detail[detail.date.isin(pd.to_datetime(["2025-02-01", "2025-05-01", "2025-08-01", "2025-11-01"]))]
    specified.to_csv(RESULT / "问题2_指定日期结果.csv", index=False, encoding="utf-8-sig")
    monthly = daily.groupby(daily.date.dt.to_period("M")).agg(
        days=("date", "size"), plan_cost_yuan=("plan_cost_yuan", "sum"),
        emergency_cost_yuan=("emergency_cost_yuan", "sum"),
        total_cost_yuan=("total_cost_yuan", "sum"), emergency_kwh=("emergency_kwh", "sum"),
        surplus_kwh=("surplus_kwh", "sum"),
    ).reset_index()
    monthly["month"] = monthly.date.astype(str); monthly.drop(columns="date").to_csv(RESULT / "问题2_月度汇总.csv", index=False, encoding="utf-8-sig")
    figure_paths = create_figures(detail, daily, forecast)

    manifest_path = RESULT / "问题2_复现清单.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["unique_reproduction_command"] = "python 03_代码/问题2_完整复现.py"
    manifest["reproduction_driver_sha256"] = sha256(ROOT / "03_代码" / "问题2_完整复现.py")
    manifest["postprocess_code_sha256"] = sha256(Path(__file__))
    manifest["validation"] = report
    extra = [RESULT / "问题2_总体验证.json", RESULT / "问题2_指定日期结果.csv", RESULT / "问题2_月度汇总.csv"]
    manifest["postprocess_outputs"] = {str(p.relative_to(ROOT)): sha256(p) for p in extra}
    manifest["figures"] = {str(Path(p).relative_to(ROOT)): sha256(Path(p)) for p in figure_paths if Path(p).exists()}
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"validation": report["status"], "figures": len(figure_paths), "total_cost_yuan": report["metrics"]["total_cost_yuan"]}, ensure_ascii=False))
    return 0 if report["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
