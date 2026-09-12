"""把第三问推荐方案 H48 写入官方 result3.xlsx 模板并进行结构校验。"""
from pathlib import Path
import json

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
RESULT_DIR = ROOT / "04_结果" / "问题3_10场景完整求解"
OUTPUT = ROOT / "04_结果" / "result3_问题3_结果.xlsx"
DT = 1 / 6


def interval_label(t: int) -> str:
    start = t * 10
    end = (t + 1) * 10
    sh, sm = divmod(start, 60)
    eh, em = divmod(end, 60)
    return f"{sh}:{sm:02d}-{eh % 24}:{em:02d}" + ("+1" if end >= 1440 else "")


def emergency_episodes(part: pd.DataFrame) -> list[tuple[str, float]]:
    active = part.emergency_kwh.to_numpy(float) > 1e-8
    values = part.emergency_kwh.to_numpy(float)
    episodes = []
    start = None
    for t, flag in enumerate(np.r_[active, False]):
        if flag and start is None:
            start = t
        elif not flag and start is not None:
            episodes.append((f"{interval_label(start).split('-')[0]}-{interval_label(t-1).split('-')[1]}",
                             float(values[start:t].sum())))
            start = None
    return episodes


def write_matrix(ws, detail: pd.DataFrame, value_col: str, cost_col: str):
    for t in range(144):
        ws.cell(1, t + 2).value = interval_label(t)
    for row, (date, part) in enumerate(detail.groupby("date", sort=True), 2):
        part = part.sort_values("t")
        ws.cell(row, 1).value = pd.Timestamp(date).to_pydatetime()
        for t, value in enumerate(part[value_col].to_numpy(float), 2):
            ws.cell(row, t).value = float(value)
        ws.cell(row, 146).value = float(part[value_col].sum())
        ws.cell(row, 147).value = float(part[cost_col].iloc[0])


def main():
    detail = pd.read_csv(RESULT_DIR / "问题3_逐时段结果.csv", parse_dates=["date"])
    daily = pd.read_csv(RESULT_DIR / "问题3_每日指标.csv", parse_dates=["date"])
    summary = pd.read_csv(RESULT_DIR / "问题3_两方案汇总.csv")
    validation = json.loads((RESULT_DIR / "问题3_总体验证.json").read_text(encoding="utf-8"))
    detail = detail.loc[detail.policy.eq("H48")].sort_values(["date", "t"]).copy()
    daily = daily.loc[daily.policy.eq("H48")].sort_values("date").copy()
    assert len(detail) == 334 * 144 and len(daily) == 334
    assert detail.groupby("date").size().eq(144).all()

    detail = detail.merge(
        daily[["date", "baseline_cost_yuan", "adjustment_cost_yuan"]], on="date", how="left",
        validate="many_to_one"
    )
    detail["adjusted_regular_cost_yuan"] = detail.baseline_cost_yuan + detail.adjustment_cost_yuan
    template = next(ROOT.rglob("result3.xlsx"))
    wb = openpyxl.load_workbook(template)
    plan_ws, adjust_ws, charge_ws, emergency_ws = [wb[s] for s in wb.sheetnames]
    write_matrix(plan_ws, detail, "baseline_grid_kwh", "baseline_cost_yuan")
    write_matrix(adjust_ws, detail, "committed_grid_kwh", "adjusted_regular_cost_yuan")

    if charge_ws.max_row > 1:
        charge_ws.delete_rows(2, charge_ws.max_row - 1)
    row = 2
    for date, part in detail.groupby("date", sort=True):
        part = part.sort_values("t")
        for block in range(6):
            sub = part.iloc[block * 24:(block + 1) * 24]
            charge_ws.cell(row, 1).value = pd.Timestamp(date).to_pydatetime() if block == 0 else None
            charge_ws.cell(row, 2).value = f"{block * 4}:00-{(block + 1) * 4}:00"
            charge_ws.cell(row, 3).value = float(sub.charge_exec_kwh.sum())
            charge_ws.cell(row, 4).value = float(sub.discharge_exec_kwh.sum())
            if block == 0:
                charge_ws.cell(row, 5).value = "0:00"
                charge_ws.cell(row, 6).value = float(part.iloc[0].E_start_kwh)
            elif block == 1:
                charge_ws.cell(row, 5).value = "24:00"
                charge_ws.cell(row, 6).value = float(part.iloc[-1].E_end_kwh)
            row += 1

    if emergency_ws.max_row > 1:
        emergency_ws.delete_rows(2, emergency_ws.max_row - 1)
    row = 2
    for date, part in detail.groupby("date", sort=True):
        episodes = emergency_episodes(part.sort_values("t"))
        if not episodes:
            emergency_ws.cell(row, 1).value = pd.Timestamp(date).to_pydatetime()
            emergency_ws.cell(row, 2).value = "无"
            emergency_ws.cell(row, 3).value = 0.0
            row += 1
            continue
        for i, (label, amount) in enumerate(episodes):
            emergency_ws.cell(row, 1).value = pd.Timestamp(date).to_pydatetime() if i == 0 else None
            emergency_ws.cell(row, 2).value = label
            emergency_ws.cell(row, 3).value = amount
            row += 1

    notes = wb.create_sheet("模型说明与汇总")
    notes.append(["项目", "内容"])
    rows = [
        ("推荐方案", "H48：预测48 h、每6 h滚动更新并只执行前6 h"),
        ("计划购电量", "每日0:00发布的基准购电计划"),
        ("调整购电量", "后三次滚动调整完成后的最终有效购电承诺；全天费用为基准计划费用加净调整费用"),
        ("净调整量", "逐时段审计CSV中的 adjustment_grid_kwh=最终有效承诺-0:00基准计划"),
        ("完整验证通过", str(bool(validation["passed"]))),
        ("H24总费用/元", float(summary.loc[summary.policy.eq("H24"), "total_cost_yuan"].iloc[0])),
        ("H48总费用/元", float(summary.loc[summary.policy.eq("H48"), "total_cost_yuan"].iloc[0])),
        ("H48节省/元", float(validation["h24_minus_h48_total_yuan"])),
        ("配对日均差95%CI/元", str(validation["paired_mean_ci95_yuan_per_day"])),
        ("求解模式", "10条历史整段残差情景；单次MILP时限120 s；H24/H48均无回退"),
    ]
    for item in rows:
        notes.append(item)
    notes.freeze_panes = "A2"
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 88
    for cell in notes[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(OUTPUT)

    check = openpyxl.load_workbook(OUTPUT, read_only=True, data_only=False)
    assert check[check.sheetnames[0]].max_row == 335
    assert check[check.sheetnames[1]].max_row == 335
    assert check[check.sheetnames[2]].max_row == 1 + 334 * 6
    assert check[check.sheetnames[3]].max_row > 1
    print(OUTPUT)
    print({s: (check[s].max_row, check[s].max_column) for s in check.sheetnames})


if __name__ == "__main__":
    main()
