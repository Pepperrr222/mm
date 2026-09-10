"""Render the report to self-contained HTML and an editable DOCX."""
from pathlib import Path
import base64
import re
import xml.etree.ElementTree as ET
import markdown
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
STEM='ABCD题目分析与选题建议'
source=(ROOT/(STEM+'.md')).read_text(encoding='utf-8')
fragment=markdown.markdown(source,extensions=['tables','fenced_code','toc'],output_format='xhtml')

def embed(match):
    p=ROOT/match.group(1)
    if not p.is_file():raise FileNotFoundError(p)
    return 'src="data:image/png;base64,'+base64.b64encode(p.read_bytes()).decode()+'"'

embedded=re.sub(r'src="([^\"]+\.png)"',embed,fragment)
html='''<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>A—D题内容、难点与选题建议</title><style>
:root{color-scheme:light}body{font-family:"Noto Sans CJK SC","Microsoft YaHei",sans-serif;color:#213047;
margin:0;background:#eef2f6;line-height:1.85;font-size:16px}main{max-width:1080px;margin:32px auto;
background:white;padding:44px 54px;box-shadow:0 5px 25px #182e4710;border-radius:12px}
h1{font-size:30px;color:#174671;line-height:1.4}h2{font-size:24px;border-bottom:2px solid #e2eaf1;
padding-bottom:9px;margin-top:52px;color:#174671}h3{font-size:19px;margin-top:30px;color:#315b79}
img{display:block;width:100%;height:auto;margin:24px auto;border:1px solid #edf0f2;border-radius:5px}
table{border-collapse:collapse;width:100%;font-size:14px;margin:24px 0;display:block;overflow:auto}
th{background:#edf3f9;color:#174671}td,th{padding:10px 12px;border:1px solid #dbe3ea;text-align:left;min-width:50px}
tr:nth-child(even){background:#fafbfd}pre{padding:18px 20px;background:#f3f6f9;overflow:auto;border-left:3px solid #4f82ae;
font-size:13px;line-height:1.65;white-space:pre-wrap}code{font-family:Consolas,"Noto Sans Mono CJK SC",monospace;
overflow-wrap:anywhere}a{color:#226299}strong{color:#163e60}li{margin:7px 0}
.note{font-size:14px;color:#64768b}.nav{display:flex;gap:10px;flex-wrap:wrap;margin:24px 0}
.nav a{background:#edf3f9;border-radius:5px;padding:6px 12px;text-decoration:none}
@media(max-width:700px){main{margin:0;padding:24px 18px;border-radius:0}body{font-size:15px}h1{font-size:25px}h2{font-size:21px}}
@media print{body{background:white;font-size:10pt}main{margin:0;padding:0;max-width:none;box-shadow:none}
h2,h3{break-after:avoid}img,tr,pre{break-inside:avoid}table{font-size:8pt;display:table}.nav{display:none}a{color:inherit}}
</style></head><body><main><p class="note">本地附件分析 · 本科数学专业 · 2026-09-10 · 图表已内嵌，可离线阅读</p>'''+embedded+'''</main></body></html>'''
(ROOT/(STEM+'.html')).write_text(html,encoding='utf-8')

doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.75)
sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5)
normal.element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.line_spacing=1.2
for s in ['Title','Heading 1','Heading 2','Heading 3']:
    doc.styles[s].font.name='Calibri'
    doc.styles[s].element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
    doc.styles[s].font.color.rgb=RGBColor.from_string('174671')
    doc.styles[s].paragraph_format.keep_with_next=True
footer=sec.footer.paragraphs[0]
footer.alignment=2
footer.add_run('A—D题目分析与选题建议  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def inline(par,node,bold=False,italic=False,mono=False):
    def add(text):
        if not text:return
        run=par.add_run(text);run.bold=bold;run.italic=italic
        if mono:run.font.name='Consolas';run.font.size=Pt(9)
    add(node.text)
    for child in node:
        if child.tag=='img':
            path=ROOT/child.attrib['src'];par.add_run().add_picture(str(path),width=Inches(6.5))
        elif child.tag=='br':par.add_run().add_break()
        else:
            inline(par,child,bold or child.tag=='strong',italic or child.tag=='em',mono or child.tag=='code')
            if child.tag=='a':par.add_run(' ('+child.attrib['href']+')').font.size=Pt(8)
        add(child.tail)

root=ET.fromstring('<body>'+fragment+'</body>')
for node in root:
    tag=node.tag
    if re.fullmatch('h[1-6]',tag):
        level=int(tag[1]);par=doc.add_paragraph(style='Title' if level==1 else f'Heading {min(level-1,3)}')
        inline(par,node)
    elif tag=='p':
        par=doc.add_paragraph();inline(par,node)
    elif tag in ['ul','ol']:
        for li in node:
            par=doc.add_paragraph(style='List Bullet' if tag=='ul' else 'List Number');inline(par,li)
    elif tag=='pre':
        par=doc.add_paragraph();par.paragraph_format.space_before=Pt(5)
        par.paragraph_format.line_spacing=1.05
        run=par.add_run(''.join(node.itertext()));run.font.name='Consolas';run.font.size=Pt(8.5)
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F3F6F9');par._p.get_or_add_pPr().append(shade)
    elif tag=='table':
        table_rows=node.findall('.//tr');cols=max(len(tr) for tr in table_rows)
        table=doc.add_table(rows=0,cols=cols);table.style='Table Grid'
        for ri,tr in enumerate(table_rows):
            cells=table.add_row().cells
            for ci,td in enumerate(tr):
                p=cells[ci].paragraphs[0];inline(p,td,bold=ri==0)
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.1
                for run in p.runs:run.font.size=Pt(8.5)
                if ri==0:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDF3F9');cells[ci]._tc.get_or_add_tcPr().append(sh)
            props=table.rows[-1]._tr.get_or_add_trPr()
            no_split=OxmlElement('w:cantSplit');props.append(no_split)
            if ri==0:props.append(OxmlElement('w:tblHeader'))
        doc.add_paragraph().paragraph_format.space_after=Pt(0)

doc.core_properties.title='A—D题内容、难点与选题建议'
doc.core_properties.subject='本科组数学专业；题面分析、附件统计、建模难点与选题建议'
doc.core_properties.author=''
doc.save(ROOT/(STEM+'.docx'))
print('Created',STEM+'.html','and',STEM+'.docx')
