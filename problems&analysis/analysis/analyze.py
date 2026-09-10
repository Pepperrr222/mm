"""Read-only analysis of supplied A–D data; write derived artifacts under analysis/.

Run from any directory: python analysis/analyze.py
Dependencies: analysis/requirements.txt. No official simulator is contacted.
"""
from pathlib import Path
import json
import re
import hashlib
from datetime import datetime, timedelta
import numpy as np
import pandas as pd
from openpyxl import load_workbook
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from matplotlib.patches import Circle, Polygon
from scipy.optimize import linprog, milp, Bounds, LinearConstraint
from scipy.sparse import lil_matrix

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'analysis'
FIG = OUT / 'figures'
TABLE = OUT / 'tables'
FIG.mkdir(exist_ok=True)
TABLE.mkdir(exist_ok=True)
font = Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
if font.exists():
    from matplotlib import font_manager
    font_manager.fontManager.addfont(str(font))
    plt.rcParams['font.family'] = FontProperties(fname=str(font)).get_name()
plt.rcParams.update({'axes.unicode_minus': False, 'font.size': 10, 'figure.dpi': 120,
                     'axes.spines.top': False, 'axes.spines.right': False})
COLORS = ['#2864A0', '#DA8051', '#39846B', '#9B70AB']
SUMMARY = {}

def save(name):
    plt.savefig(FIG / (name + '.png'), dpi=180, bbox_inches='tight')
    plt.savefig(FIG / (name + '.svg'), bbox_inches='tight')
    plt.close()

def rows(path, sheet=0):
    wb = load_workbook(ROOT / path, read_only=True, data_only=True)
    ws = wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]
    result = list(ws.values)
    wb.close()
    return result

def csv(data, name):
    pd.DataFrame(data).to_csv(TABLE / (name + '.csv'), index=False, encoding='utf-8-sig')

def audit():
    manifest = []
    for q in 'ABCD':
        for p in sorted((ROOT / 'problems' / (q + '题')).rglob('*')):
            if not p.is_file():
                continue
            entry = {'file': str(p.relative_to(ROOT)), 'bytes': p.stat().st_size,
                     'sha256': hashlib.sha256(p.read_bytes()).hexdigest()}
            if p.suffix == '.xlsx':
                wb = load_workbook(p, read_only=True, data_only=True)
                entry['sheets'] = [{'name': s.title, 'rows': s.max_row, 'columns': s.max_column}
                                   for s in wb]
                wb.close()
            manifest.append(entry)
    (OUT / 'input_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2))

def analyze_a():
    a = np.array(rows('problems/A题/附件/附件1.xlsx')[1:], float)
    r = np.array(rows('problems/A题/附件/附件2.xlsx')[1:], float)
    fig, axes = plt.subplots(1, 3, figsize=(15, 4.2), layout='constrained')
    axes[0].plot(a[:, 0]/3600, a[:, 1], color=COLORS[0])
    axes[0].set(xlabel='时间 / h', ylabel='烘房温度 / ℃', title='A：烘房边界温度（附件1）')
    axes[1].plot(a[:, 0]/3600, a[:, 2], color=COLORS[2])
    axes[1].set(xlabel='时间 / h', ylabel='烘房水分浓度 / (kg/kg)', title='外界浓度并非药材内部含水率')
    axes[2].plot(r[:, 0]/3600, r[:, 1], color=COLORS[1])
    axes[2].set(xlabel='时间 / h', ylabel='半径 / cm', title='A：药材收缩（附件2）')
    for ax in axes:
        ax.grid(alpha=.2)
    save('01_A_environment_shrinkage')
    alpha = .36/(820*2600)
    d0 = 7e-9*np.exp(-.89/2.55)
    s = {'environment_rows': len(a), 'environment_end_h': a[-1,0]/3600,
         'environment_step_s': sorted(np.unique(np.diff(a[:,0])).tolist()),
         'temperature_min_max': [a[:,1].min(), a[:,1].max()],
         'temperature_last_hour_mean': a[a[:,0]>=10800,1].mean(),
         'concentration_min_max': [a[:,2].min(), a[:,2].max()],
         'concentration_last_hour_mean': a[a[:,0]>=10800,2].mean(),
         'radius_rows': len(r), 'radius_end_h': r[-1,0]/3600,
         'radius_step_s': sorted(np.unique(np.diff(r[:,0])).tolist()),
         'radius_start_end_cm': [r[0,1],r[-1,1]],
         'radius_increase_steps': int(np.sum(np.diff(r[:,1])>0)),
         'radius_decline_fraction': 1-r[-1,1]/r[0,1],
         'final_area_ratio': (r[-1,1]/r[0,1])**2,
         'alpha_m2_s': alpha, 'initial_D_m2_s': d0,
         'thermal_diffusion_scale_s': .02**2/alpha,
         'moisture_diffusion_scale_h': .02**2/d0/3600,
         'Bi_heat_radial': 25*.02/.36, 'Bi_mass_initial_radial': 8e-7*.02/d0,
         'missing_numeric': int(np.isnan(a).sum()+np.isnan(r).sum())}
    SUMMARY['A'] = s
    csv({'time_s': a[:,0], 'temperature_C':a[:,1], 'ambient_concentration':a[:,2]}, 'A_environment')
    csv({'time_s':r[:,0], 'radius_cm':r[:,1]}, 'A_radius')

def analyze_c():
    day = np.array([x[1:] for x in rows('problems/C题/附件/附件1.xlsx')[1:]],float)
    raw_l = rows('problems/C题/附件/附件2.xlsx','小区负载')
    raw_p = rows('problems/C题/附件/附件2.xlsx','光伏发电实际功率')
    raw_price = rows('problems/C题/附件/附件4.xlsx')
    dates = pd.DatetimeIndex([r[0] for r in raw_l[1:]])
    assert [r[0] for r in raw_l[1:]] == [r[0] for r in raw_p[1:]] == [r[0] for r in raw_price[1:]]
    assert raw_l[0][1:] == raw_p[0][1:] == raw_price[0][1:]
    load = np.array([r[1:] for r in raw_l[1:]],float)
    pv = np.array([r[1:] for r in raw_p[1:]],float)
    price = np.array([r[1:] for r in raw_price[1:]],float)
    hours = np.arange(1,145)/6
    assert load.shape == pv.shape == price.shape == (365,144)
    assert not dates.duplicated().any()
    fig, axes = plt.subplots(2,2,figsize=(13,8),layout='constrained')
    for ax, date in zip(axes.flat, ['2025-03-20','2025-06-21','2025-09-23','2025-12-21']):
        i = dates.get_loc(date)
        ax.plot(hours,load[i]/1000,label='负载',color=COLORS[0])
        ax.plot(hours,pv[i]/1000,label='光伏实际',color=COLORS[1])
        ax.fill_between(hours,0,np.maximum(load[i]-pv[i],0)/1000,color=COLORS[0],alpha=.10,label='正净负载')
        ax.set(xlabel='时刻 / h（按右端点解释）',ylabel='功率 / MW',title=date,xlim=(0,24))
        ax.legend(fontsize=9); ax.grid(alpha=.2)
    save('03_C_four_days')
    daily = pd.DataFrame({'date':dates,'load_kWh':load.sum(axis=1)/6,'pv_kWh':pv.sum(axis=1)/6,
                          'positive_net_kWh':np.maximum(load-pv,0).sum(axis=1)/6,
                          'surplus_pv_kWh':np.maximum(pv-load,0).sum(axis=1)/6,
                          'price_mean':price.mean(axis=1)})
    daily['month'] = dates.month
    monthly = daily.groupby('month').mean(numeric_only=True).reset_index()
    fig, axes = plt.subplots(1,3,figsize=(15,4.5),layout='constrained')
    axes[0].plot(monthly.month,monthly.load_kWh/1000,'o-',label='日均负载',color=COLORS[0])
    axes[0].plot(monthly.month,monthly.pv_kWh/1000,'o-',label='日均光伏',color=COLORS[1])
    axes[0].set(xlabel='月份',ylabel='日均电量 / MWh',title='C：年度能量规模'); axes[0].legend()
    net = load-pv
    im=axes[1].imshow(net/1000,aspect='auto',origin='lower',extent=[0,24,1,366],cmap='RdBu_r',
                      vmin=-np.max(np.abs(net))/1000,vmax=np.max(np.abs(net))/1000)
    axes[1].set(xlabel='时刻 / h',ylabel='年内第几天',title='净负载：红色缺电，蓝色光伏盈余')
    fig.colorbar(im,ax=axes[1],label='MW')
    quant=np.quantile(price,[.1,.5,.9],axis=0)
    axes[2].fill_between(hours,quant[0],quant[2],alpha=.2,color=COLORS[2],label='全年10%–90%分位')
    axes[2].plot(hours,quant[1],color=COLORS[2],label='全年中位数')
    axes[2].plot(hours,day[:,0],color=COLORS[3],label='附件1固定日价',lw=1)
    axes[2].set(xlabel='时刻 / h',ylabel='元/kWh',title='C：电价波动与固定日价');axes[2].legend(fontsize=8)
    save('04_C_annual_patterns')
    # Forecast value h is paired with actual power at issue timestamp + h hours.
    # Final-day out-of-year targets are excluded, never wrapped into January.
    actual = {}
    for di,d in enumerate(dates):
        for k in range(144): actual[d+timedelta(minutes=(k+1)*10)] = pv[di,k]
    frecs=[]
    current_date=None
    for row in rows('problems/C题/附件/附件3.xlsx')[1:]:
        if row[0]: current_date=pd.Timestamp(row[0])
        issue_hour=int(str(row[1]).split(':')[0])
        issue=current_date+timedelta(hours=issue_hour)
        for lead,value in enumerate(row[2:],1):
            target=issue+timedelta(hours=lead)
            if target not in actual: continue
            frecs.append({'issue':issue,'issue_hour':issue_hour,'lead_h':lead,'target':target,
                          'forecast_kW':float(value),'actual_kW':actual[target],
                          'error_kW':float(value)-actual[target]})
    f=pd.DataFrame(frecs)
    # Use actual > 100 kW to avoid night zeros dominating the error summary.
    metrics=[]
    for ih,g in f.groupby('issue_hour'):
        light=g[g.actual_kW>100]
        metrics.append({'issue_hour':int(ih),'n':len(g),'daylight_n':len(light),
                        'MAE_all_kW':g.error_kW.abs().mean(),
                        'MAE_daylight_kW':light.error_kW.abs().mean(),
                        'RMSE_daylight_kW':np.sqrt(np.mean(light.error_kW**2)),
                        'bias_daylight_kW':light.error_kW.mean()})
    # Compare revisions on exactly the same target timestamps (13:00–24:00).
    matched=[]
    for d in dates:
        subset=f[(f.issue>=d)&(f.issue<d+timedelta(days=1)) & (f.target>d+timedelta(hours=12))
                 & (f.target<=d+timedelta(days=1)) & (f.actual_kW>100) & (f.issue_hour.isin([0,6,12]))]
        wide=subset.pivot(index='target',columns='issue_hour',values='error_kW').dropna()
        if all(h in wide.columns for h in [0,6,12]):
            for t,row in wide.iterrows():
                matched.append({'target':t,'error_00':row[0],'error_06':row[6],'error_12':row[12]})
    matched=pd.DataFrame(matched)
    matched_mae={h:float(matched['error_'+h].abs().mean()) for h in ['00','06','12']}
    fig,axes=plt.subplots(1,2,figsize=(12,4.5),layout='constrained')
    for ih,g in f[f.actual_kW>100].groupby('issue_hour'):
        z=g.groupby('lead_h').error_kW.apply(lambda v:v.abs().mean())
        axes[0].plot(z.index,z.values,label=f'{ih:02d}:00发布',alpha=.85)
    axes[0].set(xlabel='预报提前量 / h',ylabel='MAE / kW',title='按提前量分组（白天实际功率>100 kW）')
    axes[0].legend(fontsize=8)
    axes[1].bar(list(matched_mae.keys()),list(matched_mae.values()),color=COLORS[:3])
    axes[1].set(xlabel='当天预报发布时间',ylabel='MAE / kW',title='同一批当天下午目标：更新预报是否改善')
    for i,v in enumerate(matched_mae.values()): axes[1].text(i,v+4,f'{v:.1f}',ha='center')
    save('05_C_forecast_errors')
    # Diagnostic Q1 LP: 144 ten-minute intervals, terminal SOC=6000, both efficiencies .9.
    # Variables: grid[n], charge[n], discharge[n], spill[n], SOC[n+1]. All energy in kWh.
    n=144;nv=5*n+1
    objective=np.zeros(nv);objective[:n]=day[:,0]
    aeq=lil_matrix((2*n,nv));beq=np.zeros(2*n)
    for t in range(n):
        aeq[t,t]=1;aeq[t,n+t]=-1;aeq[t,2*n+t]=1;aeq[t,3*n+t]=-1
        beq[t]=(day[t,1]-day[t,2])/6
        aeq[n+t,4*n+t+1]=1;aeq[n+t,4*n+t]=-1
        aeq[n+t,n+t]=-.9;aeq[n+t,2*n+t]=1/.9
    bounds=[(0,None)]*n+[(0,5000/6)]*(2*n)+[(0,None)]*n+[(1200,10800)]*(n+1)
    bounds[4*n]=bounds[-1]=(6000,6000)
    sol=linprog(objective,A_eq=aeq.tocsr(),b_eq=beq,bounds=bounds,method='highs')
    assert sol.success,sol.message
    residual=float(np.max(np.abs(aeq@sol.x-beq)))
    simultaneous=int(np.sum((sol.x[n:2*n]>1e-6)&(sol.x[2*n:3*n]>1e-6)))
    assert residual<1e-5 and simultaneous==0
    no_store=float(np.dot(day[:,0],np.maximum(day[:,1]-day[:,2],0)/6))
    fig,axes=plt.subplots(2,1,figsize=(12,7),layout='constrained')
    axes[0].plot(hours,day[:,1]/1000,label='负载',color=COLORS[0])
    axes[0].plot(hours,day[:,2]/1000,label='光伏',color=COLORS[1])
    axes[0].step(hours,sol.x[:n]*6/1000,where='pre',label='计划购电（初步LP）',color=COLORS[2])
    axes[0].set(ylabel='功率 / MW',title='C问题1：明确假设下的可行调度示例');axes[0].legend()
    axes[1].plot(np.arange(145)/6,sol.x[4*n:]/1000,color=COLORS[3],label='储电量')
    axes[1].axhline(1.2,color='gray',ls='--');axes[1].axhline(10.8,color='gray',ls='--')
    axes[1].set(xlabel='时刻 / h',ylabel='储电量 / MWh',title=f'当日成本 {sol.fun:,.1f} 元；无储能基准 {no_store:,.1f} 元')
    save('06_C_q1_dispatch_probe')
    csv(daily,'C_daily_energy');csv(monthly,'C_monthly_energy');csv(metrics,'C_forecast_metrics')
    csv(matched,'C_forecast_matched_targets')
    csv({'interval_end_h':hours,'price':day[:,0],'load_kW':day[:,1],'pv_kW':day[:,2],
         'grid_kWh':sol.x[:n],'charge_kWh':sol.x[n:2*n],'discharge_kWh':sol.x[2*n:3*n],
         'spill_kWh':sol.x[3*n:4*n],'SOC_end_kWh':sol.x[4*n+1:]},'C_q1_LP_probe')
    SUMMARY['C']={'days':len(dates),'points_per_day':144,'annual_points_per_series':load.size,
                  'evaluation_days':int(np.sum(dates.month>=2)),
                  'missing_numeric':int(np.isnan(load).sum()+np.isnan(pv).sum()+np.isnan(price).sum()),
                  'load_min_max_kW':[load.min(),load.max()], 'pv_min_max_kW':[pv.min(),pv.max()],
                  'price_min_max':[price.min(),price.max()],
                  'annual_load_MWh':load.sum()/6000,'annual_pv_MWh':pv.sum()/6000,
                  'pv_energy_share':pv.sum()/load.sum(),
                  'surplus_intervals':int(np.sum(pv>load)),
                  'surplus_days':int(np.sum(np.any(pv>load,axis=1))),
                  'daily_load_mean_kWh':daily.load_kWh.mean(),
                  'storage_usable_to_daily_load':9600/daily.load_kWh.mean(),
                  'forecast_rows':1460,'forecast_valid_pairs':len(f),
                  'forecast_excluded_out_of_year_pairs':1460*24-len(f),
                  'forecast_metrics':metrics,'matched_afternoon_count':len(matched),
                  'matched_afternoon_MAE_kW':matched_mae,
                  'q1_probe':{'cost_yuan':float(sol.fun),'no_storage_cost_yuan':no_store,
                              'cost_saving_fraction':1-sol.fun/no_store,
                              'grid_kWh':float(sol.x[:n].sum()),'max_equality_residual':residual,
                              'simultaneous_charge_discharge_slots':simultaneous,
                              'SOC_min_max_kWh':[float(sol.x[4*n:].min()),float(sol.x[4*n:].max())]}}

def analyze_d():
    records=[]
    for row in rows('problems/D题/附件/附件1.xlsx')[1:]:
        fl,fu=map(int,re.findall(r'\d+',row[1]));ts,te=map(int,re.findall(r'\d+',row[2]))
        dur=te-ts;gap=int(row[3]);repeat=int(row[4])
        records.append({'id':row[0],'class':row[0][0],'fl':fl,'fu':fu,'start':ts,'duration':dur,
                        'gap':gap,'repeat':repeat,'intervals':[(ts+k*(dur+gap),te+k*(dur+gap)) for k in range(repeat)]})
    n=len(records);horizon=max(b for r in records for a,b in r['intervals'])
    occ=np.zeros((100,horizon),int)
    for r in records:
        for a,b in r['intervals']:occ[r['fl']:r['fu'],a:b]+=1
    adj=np.zeros((n,n),int);pairs=[];cross={};events=0
    for i in range(n):
        for j in range(i+1,n):
            a,b=records[i],records[j]
            if max(a['fl'],b['fl'])>=min(a['fu'],b['fu']):continue
            hits=[(max(s1,s2),min(e1,e2)) for s1,e1 in a['intervals'] for s2,e2 in b['intervals'] if max(s1,s2)<min(e1,e2)]
            if not hits:continue
            adj[i,j]=adj[j,i]=1;events+=len(hits)
            key=a['class']+b['class'];cross[key]=cross.get(key,0)+1
            pairs.append({'equipment_1':a['id'],'equipment_2':b['id'],'overlapping_occurrence_pairs':len(hits)})
    unseen=set(range(n));components=[]
    while unseen:
        todo=[unseen.pop()];comp=[]
        while todo:
            u=todo.pop();comp.append(u)
            neighbors=set(np.flatnonzero(adj[u]))&unseen
            unseen-=neighbors;todo.extend(neighbors)
        components.append(comp)
    csv(pairs,'D_conflict_pairs')
    csv([{'id':r['id'],'class':r['class'],'degree':int(adj[i].sum())} for i,r in enumerate(records)],'D_conflict_degrees')
    # Independent cell-based validation of every pair (half-open integer rectangles).
    masks=[]
    for r in records:
        mask=np.zeros_like(occ,dtype=bool)
        for a,b in r['intervals']:mask[r['fl']:r['fu'],a:b]=True
        masks.append(mask)
    check=sum(int(np.any(masks[i]&masks[j])) for i in range(n) for j in range(i+1,n))
    assert check==len(pairs)
    fig,axes=plt.subplots(1,2,figsize=(13,5),layout='constrained')
    im=axes[0].imshow(occ,origin='lower',aspect='auto',extent=[0,horizon,0,100],cmap='YlOrRd',vmin=0)
    axes[0].set(xlabel='时间 / Δt',ylabel='频段 / Δf',title='D：原计划资源占用重数（附件1）')
    fig.colorbar(im,ax=axes[0],label='同一时频单元占用数')
    axes[1].imshow(adj,cmap='Blues',interpolation='nearest',origin='upper')
    counts={c:sum(r['class']==c for r in records) for c in 'ABC'}
    for b in np.cumsum(list(counts.values()))[:-1]:
        axes[1].axvline(b-.5,color=COLORS[1],lw=.8);axes[1].axhline(b-.5,color=COLORS[1],lw=.8)
    axes[1].set(xlabel='装备序号（A/B/C顺序）',ylabel='装备序号',title=f'冲突矩阵：{len(pairs)} 对；并非单次冲突事件数')
    save('07_D_occupancy_conflicts')
    fig,axes=plt.subplots(1,2,figsize=(12,4.5),layout='constrained')
    axes[0].bar(list(cross),list(cross.values()),color=COLORS[0]);axes[0].set(ylabel='冲突装备对数',title='D：类别之间的冲突')
    ranked=sorted(enumerate(records),key=lambda x:adj[x[0]].sum(),reverse=True)[:15]
    axes[1].bar([r['id'] for i,r in ranked],[adj[i].sum() for i,r in ranked],color=COLORS[1])
    axes[1].tick_params(axis='x',rotation=60);axes[1].set(ylabel='冲突度数',title='冲突最多的15台装备')
    save('08_D_conflict_statistics')
    # Small exact diagnostic: maximize retained plans ONLY BY CANCELLATION, no shifts, no class priority.
    # This restricted optimum is an achievable baseline, NOT a Q2 optimum or a lower bound on Q2 cancellations.
    edge_ij=np.argwhere(np.triu(adj,1))
    mat=lil_matrix((len(edge_ij),n))
    for k,(i,j) in enumerate(edge_ij):mat[k,i]=mat[k,j]=1
    res=milp(-np.ones(n),integrality=np.ones(n),bounds=Bounds(0,1),
             constraints=LinearConstraint(mat.tocsr(),-np.inf,1),options={'time_limit':15})
    baseline={'solver_status':int(res.status),'message':res.message}
    if res.x is not None:
        kept=np.rint(res.x).astype(int);assert not np.any(adj[np.ix_(kept==1,kept==1)])
        baseline.update({'retained':int(kept.sum()),'cancelled':int(n-kept.sum()),'mip_gap':float(res.mip_gap)})
    SUMMARY['D']={'equipment':n,'class_counts':counts,'horizon':horizon,
                  'occurrences':sum(r['repeat'] for r in records),'conflicting_equipment_pairs':len(pairs),
                  'overlapping_occurrence_pairs':events,'cross_class_pairs':cross,
                  'equipment_in_conflict':int(np.sum(adj.sum(axis=1)>0)),
                  'component_sizes_desc':sorted([len(x) for x in components],reverse=True),
                  'raw_requested_area':int(occ.sum()),'occupied_union_cells':int(np.sum(occ>0)),
                  'conflict_cells':int(np.sum(occ>1)),'max_cell_multiplicity':int(occ.max()),
                  'requested_area_ratio':float(occ.sum()/(100*horizon)),
                  'class_patterns':{c:sorted(set((r['fu']-r['fl'],r['duration'],r['gap'],r['repeat']) for r in records if r['class']==c)) for c in 'ABC'},
                  'cancellation_only_baseline':baseline,
                  'pairwise_detection_independent_check':bool(check==len(pairs))}

def draw_b():
    fig,axes=plt.subplots(1,2,figsize=(12,5),layout='constrained')
    vertices=np.array([[0,0],[40,0],[20,20*np.sqrt(3)]])
    centroid=vertices.mean(axis=0)
    axes[0].add_patch(Polygon(vertices,fill=False,color=COLORS[0],lw=2))
    axes[0].add_patch(Circle(centroid,20,fill=False,color=COLORS[1],ls='--',label='直径40 m圆（示意）'))
    axes[0].add_patch(Circle(centroid,40/np.sqrt(3),fill=False,color=COLORS[2],label='最小包围圆：半径23.09 m'))
    axes[0].scatter(*vertices.T,color=COLORS[0]);axes[0].set_aspect('equal')
    axes[0].set(xlim=(-10,50),ylim=(-15,45),title='B问题1反例：区域直径40 m，半径20 m圆无法覆盖',xlabel='x / m',ylabel='y / m')
    axes[0].legend(fontsize=8,loc='lower center')
    ang=np.deg2rad(np.linspace(5,175,300));axes[1].plot(np.rad2deg(ang),1/np.sin(ang),color=COLORS[3])
    axes[1].set(xlabel='两条测向线交会角 / °',ylabel='相对几何放大因子 1/|sin θ|',
                title='仅为几何示意：近乎平行时误差显著放大',ylim=(0,12))
    axes[1].axvline(90,ls='--',color='gray');axes[1].grid(alpha=.2)
    save('02_B_geometry_illustration')
    SUMMARY['B']={'figure_type':'constructed geometric illustration, NOT simulator data',
                  'equilateral_side_m':40,'minimum_enclosing_radius_m':40/np.sqrt(3),
                  'full_scan_20_channels_initial_s':20*5+19,
                  'one_km_travel_s':1000/5,'formal_simulator_runs':0}

if __name__=='__main__':
    audit()
    for function in [analyze_a,draw_b,analyze_c,analyze_d]:
        print('Running',function.__name__,flush=True);function()
    (OUT/'summary.json').write_text(json.dumps(SUMMARY,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(SUMMARY,ensure_ascii=False,indent=2))
